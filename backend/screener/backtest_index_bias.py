"""
screener/backtest_index_bias.py

Answers the question "does our Bias reading actually predict what price
does next?" using REAL data -- every snapshot the Index Tracker has
already been logging to signal_logs/index_tracker_{NIFTY|BANKNIFTY}_
YYYY-MM-DD.xlsx since Round 14. No new data collection needed, this
reads what's already there. Now also works for commodities (crude oil)
the same way -- load_all_snapshots() and write_backtest_report() were
already generic on index_name, only backtest()/backtest_by_day() needed
a fix: they fall back to Fut (futures price) when Spot is None, since
commodities never have a separate spot/cash index to read one from.

METHOD: for every logged snapshot with a directional Bias (Bullish/
Bearish, Neutral is excluded since it makes no directional claim to
test), look ahead N minutes to the nearest later snapshot and check
whether spot actually moved the direction Bias implied. Aggregate into
a hit rate per Bias category per time horizon.

HONEST LIMITATION: this is only as good as how many days of logs have
accumulated. One day's worth of snapshots is a small, noisy sample --
treat any single-day result as a rough first look, not a verified edge.
The more days this runs, the more this number means something. This
script will tell you your actual current sample size so you're not
misled by a tiny one.

Run as a standalone script:
    cd backend
    venv\\Scripts\\activate
    python -m screener.backtest_index_bias
"""
import os
import glob
from datetime import datetime, timedelta

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

LOG_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "signal_logs")

DIRECTIONAL_BIASES = ("Bullish", "Bullish (Strong)", "Bearish", "Bearish (Strong)")


def load_all_snapshots(index_name):
    """
    Load and concatenate every daily Index Tracker log for one index,
    sorted chronologically by real datetime (combining each file's date
    -- from the filename -- with each row's Time column).

    Checks both layouts: the old flat one (signal_logs/index_tracker_*.xlsx,
    from before daily subfolders existed) and the new one (signal_logs/
    YYYY-MM-DD/index_tracker_*.xlsx), so nothing from before the folder
    reorg drops out of the backtest.
    """
    flat_pattern = os.path.join(LOG_DIR, f"index_tracker_{index_name}_*.xlsx")
    nested_pattern = os.path.join(LOG_DIR, "*", f"index_tracker_{index_name}_*.xlsx")
    all_paths = sorted(set(glob.glob(flat_pattern) + glob.glob(nested_pattern)))
    all_rows = []
    for path in all_paths:
        if "pre-update" in os.path.basename(path):
            continue  # archived schema-migration files, not part of the real series
        prefix = f"index_tracker_{index_name}_"
        date_str = os.path.basename(path)[len(prefix):-len(".xlsx")]
        try:
            wb = load_workbook(path)
            ws = wb["Snapshots"]
        except Exception as e:
            print(f"  (skipping {os.path.basename(path)}: {e})")
            continue
        headers = [c.value for c in ws[1]]
        if "Time" not in headers or "Spot" not in headers or "Bias" not in headers:
            continue  # older schema without what we need -- skip rather than guess
        for row in ws.iter_rows(min_row=2, values_only=True):
            d = dict(zip(headers, row))
            try:
                dt = datetime.strptime(f"{date_str} {d['Time']}", "%Y-%m-%d %H:%M:%S")
            except Exception:
                continue
            d["_datetime"] = dt
            all_rows.append(d)
    all_rows.sort(key=lambda r: r["_datetime"])
    return all_rows


def backtest(index_name, horizon_minutes):
    """
    For one index and one look-ahead horizon: returns
    {bias_label: {'correct': n, 'total': n, 'hit_rate': pct}}

    IMPORTANT: only counts NON-OVERLAPPING samples. The first version of
    this counted every single snapshot as an independent test -- but if
    Bias sits at the same value for an extended stretch (e.g. NIFTY
    reading "Bearish" continuously for 2 hours while price just drifts
    down the whole time), that's really ONE continuous move, not dozens
    of separate confirmations. Counting it as dozens inflates the hit
    rate for exactly the reason a flat, suspicious 100% showed up for
    BANKNIFTY Bearish(Strong) across every horizon in testing -- 36
    heavily-overlapping windows from what was likely one sustained trend,
    not 36 independent bets. After using a snapshot as a test point, the
    next eligible one has to be at least `horizon_minutes` later.
    """
    rows = load_all_snapshots(index_name)
    results = {b: {"correct": 0, "total": 0} for b in DIRECTIONAL_BIASES}
    next_eligible_time = None

    for i, row in enumerate(rows):
        bias = row.get("Bias")
        if bias not in DIRECTIONAL_BIASES:
            continue
        if next_eligible_time is not None and row["_datetime"] < next_eligible_time:
            continue  # still inside the previous sample's window -- skip, don't double-count
        # Commodities (crude oil etc.) have no separate spot/cash index --
        # Spot is always None for those rows by design. Fut (the futures
        # price) is the only real price series that exists for them, so
        # fall back to it rather than skipping every commodity row outright.
        spot_now = row.get("Spot")
        if spot_now is None:
            spot_now = row.get("Fut")
        if spot_now is None:
            continue

        target_time = row["_datetime"] + timedelta(minutes=horizon_minutes)
        future_row = next((r for r in rows[i + 1:] if r["_datetime"] >= target_time), None)
        if future_row is None:
            continue  # no snapshot far enough ahead yet (e.g. near end of day/data)
        spot_future = future_row.get("Spot")
        if spot_future is None:
            spot_future = future_row.get("Fut")
        if spot_future is None:
            continue

        predicted_up = bias.startswith("Bullish")
        actual_up = spot_future > spot_now
        results[bias]["total"] += 1
        if predicted_up == actual_up:
            results[bias]["correct"] += 1
        next_eligible_time = target_time  # don't reuse anything inside this window again

    for b in results:
        t = results[b]["total"]
        results[b]["hit_rate"] = round(100 * results[b]["correct"] / t, 1) if t else None

    return results, len(rows)


def backtest_by_day(index_name, horizon_minutes):
    """
    Same method as backtest() -- including the non-overlapping-sample
    rule explained there -- but broken out day by day instead of
    collapsed into one aggregate number across all history. Lets you
    see whether a given day was actually a good one for the Bias
    reading, rather than an overall average that can hide a lot of
    day-to-day variation.

    Returns {date_str: {bias_label: {'correct', 'total', 'hit_rate'}}},
    newest day first.
    """
    rows = load_all_snapshots(index_name)
    by_day = {}
    next_eligible_time = None

    for i, row in enumerate(rows):
        bias = row.get("Bias")
        if bias not in DIRECTIONAL_BIASES:
            continue
        if next_eligible_time is not None and row["_datetime"] < next_eligible_time:
            continue
        # Same fallback as backtest() above -- commodities never have a
        # Spot value, Fut is their only real price series.
        spot_now = row.get("Spot")
        if spot_now is None:
            spot_now = row.get("Fut")
        if spot_now is None:
            continue

        target_time = row["_datetime"] + timedelta(minutes=horizon_minutes)
        future_row = next((r for r in rows[i + 1:] if r["_datetime"] >= target_time), None)
        if future_row is None:
            continue
        spot_future = future_row.get("Spot")
        if spot_future is None:
            spot_future = future_row.get("Fut")
        if spot_future is None:
            continue

        day_str = row["_datetime"].strftime("%Y-%m-%d")
        if day_str not in by_day:
            by_day[day_str] = {b: {"correct": 0, "total": 0} for b in DIRECTIONAL_BIASES}

        predicted_up = bias.startswith("Bullish")
        actual_up = spot_future > spot_now
        by_day[day_str][bias]["total"] += 1
        if predicted_up == actual_up:
            by_day[day_str][bias]["correct"] += 1
        next_eligible_time = target_time

    for day_str, results in by_day.items():
        for b in results:
            t = results[b]["total"]
            results[b]["hit_rate"] = round(100 * results[b]["correct"] / t, 1) if t else None

    return dict(sorted(by_day.items(), reverse=True))


def _confirmation_bucket(horizons_confirming_str):
    """
    Buckets a "Horizons Confirming" value (e.g. "3/4", "0/2", "—") into
    one of three groups: 'Strong' (every horizon that had a real
    verdict agreed with Bias), 'Mixed' (some agreed, some conflicted),
    'Conflict' (every horizon that had a real verdict disagreed).

    Returns None for "—" or anything unparseable -- no real verdict to
    bucket yet, same "skip rather than guess" treatment backtest() already
    gives a missing Spot/Fut. This also correctly and silently excludes
    every row logged before Aug 21 2026 -- those files' headers never had
    a "Horizons Confirming" column at all (this feature didn't exist
    yet), so load_all_snapshots()'s dict(zip(headers, row)) never puts
    the key there, and .get() on a dict with no such key already returns
    None on its own -- nothing extra needed to handle that case.

    The denominator in "X/Y" only counts horizons that returned a real
    ✓/⚠ verdict (see _multi_horizon_confirms() in index_tracker.py) --
    flat/no-data horizons are excluded from Y already, so this ratio is
    "of the horizons with something real to say, how many agreed,"
    not diluted by ones that had nothing to say either way.
    """
    if not horizons_confirming_str or horizons_confirming_str == "—":
        return None
    try:
        num_str, denom_str = horizons_confirming_str.split("/")
        num, denom = int(num_str), int(denom_str)
    except (ValueError, AttributeError):
        return None
    if denom == 0:
        return None
    if num == denom:
        return "Strong (all horizons agree)"
    if num == 0:
        return "Conflict (all horizons disagree)"
    return "Mixed (some agree, some don't)"


def backtest_by_confirmation_level(index_name, horizon_minutes):
    """
    Same method and the same non-overlapping-sample rule as backtest()
    above -- but this answers a DIFFERENT question: not "does Bias
    predict price," but "does the Aug 20 multi-horizon confirmation
    feature (Confirms 5min/30min/60min -> Horizons Confirming) actually
    separate the good Bias reads from the noisy ones?" If 'Strong'
    confirmation rows hit meaningfully higher than 'Conflict' rows, the
    feature is doing real work. If they land about the same, it isn't
    adding predictive value yet regardless of how it looks on screen --
    that's a real, useful, possibly disappointing answer either way, and
    the whole point of building this rather than trusting the feature
    on faith.

    Only counts rows with a real confirmation verdict (see
    _confirmation_bucket) -- meaning this will have a noticeably smaller
    sample than backtest()'s total, since every snapshot logged before
    Aug 21 2026 has nothing to bucket here at all.

    Returns {(bias_label, confirmation_bucket): {'correct', 'total', 'hit_rate'}}
    """
    rows = load_all_snapshots(index_name)
    results = {}
    next_eligible_time = None

    for i, row in enumerate(rows):
        bias = row.get("Bias")
        if bias not in DIRECTIONAL_BIASES:
            continue
        if next_eligible_time is not None and row["_datetime"] < next_eligible_time:
            continue
        bucket = _confirmation_bucket(row.get("Horizons Confirming"))
        if bucket is None:
            continue  # no real confirmation verdict yet for this row -- can't bucket it, don't guess

        spot_now = row.get("Spot")
        if spot_now is None:
            spot_now = row.get("Fut")
        if spot_now is None:
            continue

        target_time = row["_datetime"] + timedelta(minutes=horizon_minutes)
        future_row = next((r for r in rows[i + 1:] if r["_datetime"] >= target_time), None)
        if future_row is None:
            continue
        spot_future = future_row.get("Spot")
        if spot_future is None:
            spot_future = future_row.get("Fut")
        if spot_future is None:
            continue

        key = (bias, bucket)
        if key not in results:
            results[key] = {"correct": 0, "total": 0}

        predicted_up = bias.startswith("Bullish")
        actual_up = spot_future > spot_now
        results[key]["total"] += 1
        if predicted_up == actual_up:
            results[key]["correct"] += 1
        next_eligible_time = target_time

    for key in results:
        t = results[key]["total"]
        results[key]["hit_rate"] = round(100 * results[key]["correct"] / t, 1) if t else None

    return results


def write_backtest_report(index_name, horizons=(15, 30, 60)):
    """
    Builds a day-wise backtest Excel report for one index, one row per
    (date, bias) combination, columns for each horizon -- and saves it
    to signal_logs/backtest_reports/backtest_<INDEX>_<today>.xlsx
    (regenerated fresh each time this is called, not appended to, since
    it's a computed report rather than a log of events). Returns the
    path.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    by_day_per_horizon = {h: backtest_by_day(index_name, h) for h in horizons}
    all_days = sorted({d for h in horizons for d in by_day_per_horizon[h]}, reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Backtest"

    headers = ["Date", "Bias"]
    for h in horizons:
        headers += [f"{h}min Hit%", f"{h}min Samples"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")

    for day in all_days:
        for bias in DIRECTIONAL_BIASES:
            row = [day, bias]
            has_any_sample = False
            for h in horizons:
                r = by_day_per_horizon[h].get(day, {}).get(bias, {"correct": 0, "total": 0, "hit_rate": None})
                if r["total"] > 0:
                    has_any_sample = True
                row += [r["hit_rate"] if r["hit_rate"] is not None else "—", r["total"]]
            if has_any_sample:
                ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(c.value)) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 2, 10)

    out_dir = os.path.join(LOG_DIR, "backtest_reports")
    os.makedirs(out_dir, exist_ok=True)
    today = datetime.now().strftime("%Y-%m-%d")
    path = os.path.join(out_dir, f"backtest_{index_name}_{today}.xlsx")
    wb.save(path)
    return path


def write_confirmation_report(index_name, horizons=(15, 30, 60)):
    """
    Same shape as write_backtest_report() above, but for the multi-
    horizon confirmation breakdown -- one row per (Bias, confirmation
    level), columns for each look-ahead horizon. Saved separately to
    signal_logs/backtest_reports/confirmation_backtest_<INDEX>_<today>.xlsx
    so it doesn't disturb the existing, already-trusted backtest report.
    Returns the path.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    results_per_horizon = {h: backtest_by_confirmation_level(index_name, h) for h in horizons}
    all_keys = sorted({k for h in horizons for k in results_per_horizon[h]})

    wb = Workbook()
    ws = wb.active
    ws.title = "Confirmation Backtest"

    headers = ["Bias", "Confirmation Level"]
    for h in horizons:
        headers += [f"{h}min Hit%", f"{h}min Samples"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")

    for key in all_keys:
        bias, bucket = key
        row = [bias, bucket]
        has_any_sample = False
        for h in horizons:
            r = results_per_horizon[h].get(key, {"correct": 0, "total": 0, "hit_rate": None})
            if r["total"] > 0:
                has_any_sample = True
            row += [r["hit_rate"] if r["hit_rate"] is not None else "—", r["total"]]
        if has_any_sample:
            ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(c.value)) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 2, 10)

    out_dir = os.path.join(LOG_DIR, "backtest_reports")
    os.makedirs(out_dir, exist_ok=True)
    today = datetime.now().strftime("%Y-%m-%d")
    path = os.path.join(out_dir, f"confirmation_backtest_{index_name}_{today}.xlsx")
    wb.save(path)
    return path


def print_confirmation_report(index_name, horizons=(15, 30, 60)):
    print(f"\n{'=' * 60}")
    print(f"  {index_name} — does multi-horizon confirmation actually help?")
    print(f"{'=' * 60}")
    print("(Segments the same Bias hit-rate test above by how many of the")
    print(" Aug 20 confirmation horizons agreed at the time. Strong should")
    print(" beat Conflict if this is adding real signal, not just noise.")
    print(" Only counts snapshots logged Aug 21 2026 onward -- the")
    print(" confirmation columns didn't exist before that.)")
    print()

    for horizon in horizons:
        results = backtest_by_confirmation_level(index_name, horizon)
        total_bucketed = sum(r["total"] for r in results.values())
        print(f"--- Look-ahead: {horizon} minutes  ({total_bucketed} bucketed samples) ---")
        if not results:
            print("  no bucketed samples yet")
            print()
            continue
        for (bias, bucket), r in sorted(results.items()):
            if r["total"] == 0:
                continue
            print(f"  {bias:18s} | {bucket:28s}: {r['hit_rate']:5.1f}% correct  ({r['correct']}/{r['total']} samples)")
        print()


def print_report(index_name, horizons=(15, 30, 60)):
    print(f"\n{'=' * 60}")
    print(f"  {index_name} — Bias backtest")
    print(f"{'=' * 60}")

    _, total_snapshots = backtest(index_name, horizons[0])
    print(f"Total logged snapshots found: {total_snapshots}")
    if total_snapshots < 200:
        print("(That's roughly under a day's worth at a 90s scan interval --")
        print(" treat any hit rate below as a rough first look, not a verified edge.")
        print(" Let this run for more days before trusting it.)")
    print()

    for horizon in horizons:
        results, _ = backtest(index_name, horizon)
        print(f"--- Look-ahead: {horizon} minutes ---")
        for bias, r in results.items():
            if r["total"] == 0:
                print(f"  {bias:18s}: no samples yet")
            else:
                print(f"  {bias:18s}: {r['hit_rate']:5.1f}% correct  ({r['correct']}/{r['total']} samples)")
        print()


if __name__ == "__main__":
    if load_workbook is None:
        print("openpyxl not installed -- pip install openpyxl")
    else:
        # Aug 20 2026: Gold/Silver joined the tracked instruments tonight
        # (see index_tracker.py) -- this script already worked for any
        # index_name generically, so just extending the list run here
        # rather than anything about the method changing.
        instruments = ("NIFTY", "BANKNIFTY", "CRUDEOIL", "CRUDEOILM", "GOLD", "GOLDM", "SILVER", "SILVERM")

        for idx in instruments:
            print_report(idx)

        print("\n" + "#" * 60)
        print("# Multi-horizon confirmation breakdown -- Aug 21 2026+ data only")
        print("#" * 60)
        for idx in instruments:
            print_confirmation_report(idx)