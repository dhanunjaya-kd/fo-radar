"""
Persistent next-day watchlist backtest ledger.

Every completed Next Day scan records the ranked top-N picks for that scan
session. On later scans, the same workbook is updated with the first actual
trading-day candle after each pick date. This is deliberately based on the
scanner's existing daily candles -- no extra Fyers polling is introduced.

The workbook is a research ledger, not a claim of live trading performance.
It records what the scanner selected first and what the stock actually did
on the following trading day.
"""
import os
from datetime import datetime

BACKTEST_FILE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "signal_logs", "next_day_watchlist_backtest.xlsx"
)

HEADERS = [
    "Scan Date", "Rank", "Symbol", "Sector", "Score", "EOD Price",
    "Trend Status", "Volume Status", "Sector Strength",
    "Next Trading Date", "Next Day Open", "Next Day High", "Next Day Low",
    "Next Day Close", "Close Return %", "Max Gain %", "Max Drawdown %",
    "Direction",
]


def _as_date(value):
    if value is None:
        return None
    if hasattr(value, "date") and not isinstance(value, str):
        return value.date()
    try:
        return datetime.fromisoformat(str(value)).date()
    except (TypeError, ValueError):
        try:
            return datetime.strptime(str(value), "%Y-%m-%d").date()
        except (TypeError, ValueError):
            return None


def _next_candle(candles, scan_date):
    dates = []
    for candle in candles or []:
        try:
            d = _as_date(candle.get("date"))
        except AttributeError:
            continue
        if d and scan_date and d > scan_date:
            dates.append((d, candle))
    return min(dates, key=lambda x: x[0])[1] if dates else None


def _load_or_create():
    from openpyxl import Workbook, load_workbook
    os.makedirs(os.path.dirname(BACKTEST_FILE), exist_ok=True)
    if os.path.exists(BACKTEST_FILE):
        wb = load_workbook(BACKTEST_FILE)
        ws = wb["Daily Picks"] if "Daily Picks" in wb.sheetnames else wb.create_sheet("Daily Picks")
        if ws.max_row == 1 and ws.cell(1, 1).value is None:
            ws.append(HEADERS)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Daily Picks"
        ws.append(HEADERS)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:R1"
        widths = (13, 8, 16, 18, 9, 13, 15, 20, 16, 18, 13, 13, 13, 13, 15, 14, 17, 12)
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width
    return wb, ws


def _rows_by_key(ws):
    out = {}
    for row in range(2, ws.max_row + 1):
        scan_date = ws.cell(row, 1).value
        symbol = ws.cell(row, 3).value
        if scan_date and symbol:
            out[(str(scan_date)[:10], str(symbol))] = row
    return out


def _apply_next_day_result(ws, row, candle, eod_price):
    if not candle or not eod_price:
        return False
    try:
        open_px = candle.get("open")
        high_px = candle.get("high")
        low_px = candle.get("low")
        close_px = candle.get("close")
        if any(v is None for v in (open_px, high_px, low_px, close_px)):
            return False
        scan_date = _as_date(ws.cell(row, 1).value)
        next_date = _as_date(candle.get("date"))
        if not scan_date or not next_date or next_date <= scan_date:
            return False
        close_return = (close_px - eod_price) / eod_price * 100
        max_gain = (high_px - eod_price) / eod_price * 100
        max_drawdown = (low_px - eod_price) / eod_price * 100
        values = {
            10: next_date.isoformat(), 11: open_px, 12: high_px, 13: low_px,
            14: close_px, 15: round(close_return, 2), 16: round(max_gain, 2),
            17: round(max_drawdown, 2),
            18: "UP" if close_return > 0 else "DOWN" if close_return < 0 else "FLAT",
        }
        for col, value in values.items():
            ws.cell(row, col).value = value
        return True
    except (TypeError, ValueError, ZeroDivisionError):
        return False


def record_and_resolve(raw_data, watchlist, scan_date=None):
    """
    Record today's top picks and resolve older unresolved picks whenever
    their first following daily candle is present in raw_data.

    This function never raises into the scanner: Excel tracking is
    secondary research output and must not break ranking or live scanning.
    """
    try:
        from openpyxl.styles import Font, Alignment

        wb, ws = _load_or_create()
        rows = _rows_by_key(ws)
        scan_date = _as_date(scan_date or datetime.now().date())
        if not scan_date:
            return {"recorded": 0, "resolved": 0}

        recorded = 0
        for pick in watchlist or []:
            symbol = pick.get("symbol")
            if not symbol:
                continue
            key = (scan_date.isoformat(), str(symbol))
            row = rows.get(key)
            values = [
                scan_date.isoformat(), pick.get("rank"), symbol,
                pick.get("sector"), pick.get("score"), pick.get("eod_price"),
                pick.get("trend_status"), pick.get("volume_status"),
                pick.get("sector_strength"), None, None, None, None, None,
                None, None, None, None,
            ]
            if row is None:
                ws.append(values)
                row = ws.max_row
                rows[key] = row
                recorded += 1
            else:
                # Same-day re-run refreshes the selection fields only.
                # Already resolved next-day values are never erased.
                for col in range(2, 10):
                    ws.cell(row, col).value = values[col - 1]

        resolved = 0
        for row in range(2, ws.max_row + 1):
            if ws.cell(row, 10).value:
                continue
            row_scan_date = _as_date(ws.cell(row, 1).value)
            symbol = ws.cell(row, 3).value
            eod_price = ws.cell(row, 6).value
            entry = raw_data.get(symbol) if isinstance(raw_data, dict) else None
            candle = _next_candle((entry or {}).get("daily_candles") or [], row_scan_date) if row_scan_date and entry else None
            if candle and _apply_next_day_result(ws, row, candle, eod_price):
                resolved += 1

        if "Summary" in wb.sheetnames:
            summary = wb["Summary"]
            summary.delete_rows(1, summary.max_row)
        else:
            summary = wb.create_sheet("Summary")
        summary.append([
            "Scan Date", "Picks", "Resolved", "Positive Close", "Negative Close",
            "Avg Close Return %", "Best Close %", "Worst Close %", "Close >= +5%", "Close <= -5%",
        ])
        summary.freeze_panes = "A2"
        grouped = {}
        for row in range(2, ws.max_row + 1):
            d = ws.cell(row, 1).value
            if not d:
                continue
            vals = grouped.setdefault(str(d)[:10], [])
            ret = ws.cell(row, 15).value
            vals.append(ret if isinstance(ret, (int, float)) else None)
        for d in sorted(grouped):
            vals = [v for v in grouped[d] if v is not None]
            summary.append([
                d, len(grouped[d]), len(vals), sum(v > 0 for v in vals), sum(v < 0 for v in vals),
                round(sum(vals) / len(vals), 2) if vals else None,
                round(max(vals), 2) if vals else None,
                round(min(vals), 2) if vals else None,
                sum(v >= 5 for v in vals), sum(v <= -5 for v in vals),
            ])
        for col, width in enumerate((13, 10, 10, 16, 17, 19, 15, 16, 15, 16), 1):
            summary.column_dimensions[chr(64 + col)].width = width

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for cell in summary[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        if ws.max_row >= 2:
            ws.auto_filter.ref = f"A1:R{ws.max_row}"

        tmp = BACKTEST_FILE + ".tmp"
        wb.save(tmp)
        os.replace(tmp, BACKTEST_FILE)
        return {"recorded": recorded, "resolved": resolved, "file": BACKTEST_FILE}
    except Exception as exc:
        print(f"[NextDayBacktest] workbook update skipped: {exc}")
        return {"recorded": 0, "resolved": 0, "error": str(exc)}
