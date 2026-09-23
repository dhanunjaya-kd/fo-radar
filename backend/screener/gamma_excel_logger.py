"""
screener/gamma_excel_logger.py

Durable, daily logging for Gamma Blast microstructure triggers --
mirrors this project's existing excel_logger.py conventions (same
signal_logs/ directory, same daily-file pattern, same corruption-
recovery approach, same cross-process file lock, duplicated rather
than imported per this project's own established preference for
small duplicated logic over a shared-utils dependency -- see
excel_logger.py's own _FileLock docstring for that precedent).

Purpose: the microstructure daemon (gamma_microstructure.py) already
tracks each trigger's live outcome (MFE/MAE, T1/T2/SL hit, realized R)
correctly -- but only in memory. Restart the backend and it's gone.
This file is what makes a month-long observation period actually
produce something reviewable afterward: one row per real trigger,
outcome synced every cycle, and a rebuild-on-restart path so a
mid-observation restart doesn't silently drop open positions from the
record.

Deliberately NOT re-implementing SL/target comparison here -- the
daemon already computes status/mfe_pct/mae_pct/realized_r correctly
every tick; this only PERSISTS whatever the daemon's current state
already says, same "single source of truth" reasoning the main
system's check_outcomes() already follows for its own signals.
"""
import os
import threading
import time
import zipfile
from datetime import datetime


class _FileLock:
    """Cross-process file lock, same mechanism as excel_logger.py's own
    copy (os.O_CREAT | os.O_EXCL atomicity, stale-lock takeover)."""
    def __init__(self, target_path, timeout=10, poll_interval=0.05):
        self.lock_path = target_path + ".lock"
        self.timeout = timeout
        self.poll_interval = poll_interval
        self._fd = None

    def __enter__(self):
        deadline = time.monotonic() + self.timeout
        while True:
            try:
                self._fd = os.open(self.lock_path, os.O_CREAT | os.O_EXCL | os.O_RDWR)
                return self
            except FileExistsError:
                if time.monotonic() >= deadline:
                    try:
                        stale_age = time.time() - os.path.getmtime(self.lock_path)
                    except OSError:
                        stale_age = None
                    if stale_age is not None and stale_age > self.timeout:
                        try:
                            os.remove(self.lock_path)
                        except OSError:
                            pass
                        continue
                    raise TimeoutError(f"Could not acquire lock on {self.lock_path} within {self.timeout}s.")
                time.sleep(self.poll_interval)

    def __exit__(self, exc_type, exc_val, exc_tb):
        if self._fd is not None:
            os.close(self._fd)
        try:
            os.remove(self.lock_path)
        except OSError:
            pass
        return False


try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

LOG_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "signal_logs")

COLUMNS = [
    "Timestamp", "Alert ID", "Symbol", "Option Type", "Strike", "Expiry", "DTE", "Contract",
    "Entry Premium", "Stop Loss", "Target 1", "Target 2", "Lot Size", "Risk:Reward",
    "Spot At Trigger", "OI Drop % (Phase 1)", "Volume Expansion x (Phase 3)", "Price Lift % (Phase 4)",
    "Trigger Candle", "Status", "Highest LTP", "Lowest LTP", "MFE %", "MAE %",
    "Realized R", "Exit Price", "Exit Time",
]

_lock = threading.Lock()
_TERMINAL_STATUSES = {"TARGET_2_HIT", "STOPPED_OUT", "TARGET_1_HIT_TRAILED"}


def _today_path():
    today = datetime.now().strftime("%Y-%m-%d")
    day_dir = os.path.join(LOG_DIR, today)
    os.makedirs(day_dir, exist_ok=True)
    return os.path.join(day_dir, f"gamma_blast_{today}.xlsx"), today


def _get_workbook(path):
    """Same corruption-recovery shape as excel_logger.py's own
    _get_workbook -- a BadZipFile (two processes writing at once with
    no lock, or an interrupted save) archives the bad file aside and
    starts fresh rather than failing silently for the rest of the
    day, every day, forever."""
    if os.path.exists(path):
        try:
            wb = load_workbook(path)
            ws = wb["Gamma Signals"]
            if [c.value for c in ws[1]] == COLUMNS:
                return wb
            # column layout changed -- archive old data, start fresh
            archive_path = path.replace(".xlsx", "_pre-update.xlsx")
            if not os.path.exists(archive_path):
                wb.save(archive_path)
        except zipfile.BadZipFile as e:
            archive_path = path.replace(".xlsx", f"_corrupted_{datetime.now().strftime('%H%M%S')}.xlsx")
            try:
                os.rename(path, archive_path)
                print(f"[GammaExcelLog] {os.path.basename(path)} corrupted ({e}) -- archived, starting fresh.")
            except OSError:
                pass
    wb = Workbook()
    ws = wb.active
    ws.title = "Gamma Signals"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    return wb


def log_new_trigger(trigger):
    """Appends ONE row for a brand-new alert (called once, when
    _gamma_feed_microstructure_and_alert sees a trigger it hasn't
    logged before)."""
    if not OPENPYXL_AVAILABLE:
        return
    path, _ = _today_path()
    metrics = trigger.get("metrics", {})
    with _lock, _FileLock(path):
        wb = _get_workbook(path)
        ws = wb["Gamma Signals"]
        ws.append([
            trigger.get("timestamp_ist", datetime.now().strftime("%Y-%m-%d %H:%M:%S IST")),
            trigger.get("alert_id"), trigger.get("symbol"), trigger.get("option_type"),
            trigger.get("strike"), trigger.get("expiry"), trigger.get("dte"), trigger.get("contract"),
            trigger.get("entry_price"), trigger.get("stop_loss"), trigger.get("target_1"), trigger.get("target_2"),
            trigger.get("lot_size"), trigger.get("risk_reward"), trigger.get("spot_cmp"),
            metrics.get("oi_drop_prior_pct"), metrics.get("volume_expansion_ratio"), metrics.get("price_lift_from_base_pct"),
            trigger.get("trigger_candle"), trigger.get("status", "ACTIVE"),
            trigger.get("highest_ltp"), trigger.get("lowest_ltp"), trigger.get("mfe_pct"), trigger.get("mae_pct"),
            trigger.get("realized_r"), trigger.get("exit_price"), trigger.get("exit_time_ist"),
        ])
        wb.save(path)


def sync_alert_outcomes(alerts_today):
    """
    Called every cycle with the microstructure daemon's CURRENT
    alerts_emitted list -- updates each existing row's Status/
    Highest/Lowest/MFE/MAE/Realized R/Exit in place, by matching on
    Alert ID. Does NOT re-derive outcomes itself; the daemon already
    did that (_update_alert_lifecycles), this only persists it.
    """
    if not OPENPYXL_AVAILABLE or not alerts_today:
        return
    path, _ = _today_path()
    if not os.path.exists(path):
        return  # nothing logged yet today -- log_new_trigger creates it first
    by_id = {a["alert_id"]: a for a in alerts_today}
    with _lock, _FileLock(path):
        try:
            wb = load_workbook(path)
        except Exception as e:
            print(f"[GammaExcelLog] sync failed to open {path}: {e}")
            return
        ws = wb["Gamma Signals"]
        headers = [c.value for c in ws[1]]
        if headers != COLUMNS:
            return
        col = {name: i + 1 for i, name in enumerate(headers)}
        changed = False
        for row_num in range(2, ws.max_row + 1):
            alert_id = ws.cell(row=row_num, column=col["Alert ID"]).value
            a = by_id.get(alert_id)
            if not a:
                continue
            ws.cell(row=row_num, column=col["Status"]).value = a.get("status")
            ws.cell(row=row_num, column=col["Highest LTP"]).value = a.get("highest_ltp")
            ws.cell(row=row_num, column=col["Lowest LTP"]).value = a.get("lowest_ltp")
            ws.cell(row=row_num, column=col["MFE %"]).value = a.get("mfe_pct")
            ws.cell(row=row_num, column=col["MAE %"]).value = a.get("mae_pct")
            ws.cell(row=row_num, column=col["Realized R"]).value = a.get("realized_r")
            ws.cell(row=row_num, column=col["Exit Price"]).value = a.get("exit_price")
            ws.cell(row=row_num, column=col["Exit Time"]).value = a.get("exit_time_ist")
            changed = True
        if changed:
            wb.save(path)


def rebuild_open_alerts_from_log():
    """
    Called once, when the microstructure daemon is first constructed
    (i.e. right after a restart) -- reads TODAY's log (if any) and
    returns alert-payload-shaped dicts for every row that isn't in a
    terminal status yet, so daemon.alerts_emitted can be pre-populated
    and _update_alert_lifecycles keeps tracking them against new
    ticks, instead of silently losing every open position's outcome
    tracking on every restart. Mirrors excel_logger.py's own
    _ensure_fresh() rebuild-on-restart philosophy.
    """
    if not OPENPYXL_AVAILABLE:
        return []
    path, _ = _today_path()
    if not os.path.exists(path):
        return []
    try:
        wb = load_workbook(path)
        ws = wb["Gamma Signals"]
        headers = [c.value for c in ws[1]]
        if headers != COLUMNS:
            return []
        col = {name: i + 1 for i, name in enumerate(headers)}
        rebuilt = []
        for row_num in range(2, ws.max_row + 1):
            status = ws.cell(row=row_num, column=col["Status"]).value or "ACTIVE"
            alert_id = ws.cell(row=row_num, column=col["Alert ID"]).value
            if not alert_id or status in _TERMINAL_STATUSES:
                continue
            rebuilt.append({
                "alert_id": alert_id, "strategy": "Gamma_Blast_Options_strategy",
                "symbol": ws.cell(row=row_num, column=col["Symbol"]).value,
                "contract": ws.cell(row=row_num, column=col["Contract"]).value,
                "option_type": ws.cell(row=row_num, column=col["Option Type"]).value,
                "strike": ws.cell(row=row_num, column=col["Strike"]).value,
                "expiry": ws.cell(row=row_num, column=col["Expiry"]).value,
                "lot_size": ws.cell(row=row_num, column=col["Lot Size"]).value,
                "entry_price": ws.cell(row=row_num, column=col["Entry Premium"]).value,
                "stop_loss": ws.cell(row=row_num, column=col["Stop Loss"]).value,
                "trailing_sl": ws.cell(row=row_num, column=col["Stop Loss"]).value,
                "target_1": ws.cell(row=row_num, column=col["Target 1"]).value,
                "target_2": ws.cell(row=row_num, column=col["Target 2"]).value,
                "risk_reward": ws.cell(row=row_num, column=col["Risk:Reward"]).value,
                "status": status,
                "highest_ltp": ws.cell(row=row_num, column=col["Highest LTP"]).value or ws.cell(row=row_num, column=col["Entry Premium"]).value,
                "lowest_ltp": ws.cell(row=row_num, column=col["Lowest LTP"]).value or ws.cell(row=row_num, column=col["Entry Premium"]).value,
                "mfe_pct": ws.cell(row=row_num, column=col["MFE %"]).value or 0.0,
                "mae_pct": ws.cell(row=row_num, column=col["MAE %"]).value or 0.0,
                "realized_r": ws.cell(row=row_num, column=col["Realized R"]).value or 0.0,
                "exit_price": ws.cell(row=row_num, column=col["Exit Price"]).value,
                "exit_time_ist": ws.cell(row=row_num, column=col["Exit Time"]).value,
                "timestamp_ist": ws.cell(row=row_num, column=col["Timestamp"]).value,
            })
        if rebuilt:
            print(f"[GammaExcelLog] Rebuilt {len(rebuilt)} still-open alert(s) from {os.path.basename(path)} -- restart resumes tracking instead of losing them.")
        return rebuilt
    except Exception as e:
        print(f"[GammaExcelLog] rebuild failed: {e}")
        return []
