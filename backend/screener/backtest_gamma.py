"""
screener/backtest_gamma.py

Oct 3 2026: the Gamma Blast Options strategy gets its OWN backtest report -- same engine, same PDF
format as the Sniper Signals report, but a separate file built only from Gamma's own trade log.

WHERE THE TRADES COME FROM
--------------------------
Two files, both already written by the live Gamma code (nothing new is logged here):
  * signal_logs/gamma_strategy_signals_clean.xlsx   (gamma_trade_tracker.py) -- one row per option
    CONTRACT, lifecycle tracked across days, including overnight carry-forwards. This is the primary
    source: it is the only file that keeps updating a trade after the day it was triggered on.
  * signal_logs/<date>/gamma_blast_<date>.xlsx      (gamma_excel_logger.py) -- one row per trigger,
    with the lot size and the three trigger metrics the tracker does not carry. Used to fill those in,
    and as the fallback source for any contract the tracker does not have.
Same contract in both -> one trade (the tracker keeps one trade per contract, so does this).

HOW A TRADE BECOMES A P&L NUMBER (mirrors gamma_microstructure.py, nothing re-invented)
---------------------------------------------------------------------------------------
Entry = the logged premium. SL = 75% of entry (so 1R = 25% of the premium). T1 = entry + 1.6R,
T2 = entry + 2.8R. The live engine books HALF at T1 and moves the stop on the other half to entry:
  STOPPED_OUT           -> -1.0R
  TARGET_2_HIT, no T1   ->  2.8R  (jumped straight through both targets)
  TARGET_2_HIT after T1 ->  0.5 x 1.6R + 0.5 x 2.8R = 2.2R
  TARGET_1_HIT_TRAILED  ->  0.5 x 1.6R + 0.5 x 0R   = 0.8R
The logged Realized R is used as is. The engine only stores the LAST leg's price as "Exit Price", which
would misstate a two-leg trade, so the P&L here uses the blended exit price = entry + R x risk -- the same
number the engine's own Realized R already implies. P&L = lot size x (blended exit - entry), one lot.

WHAT IS NEVER GUESSED (same rule as the Sniper report): a trade still open (ACTIVE / TARGET_1_HIT) is not
counted as a win, loss or flat; a row with no lot size, no exit time, or an R that cannot be recovered
unambiguously is excluded and counted, not filled in.

LIMITS worth stating next to any number this produces: exits are the engine's own LEVELS (an option that
gaps through a stop would really fill worse), no brokerage / slippage / taxes, and a month or two of
triggers is a small sample -- the PDF's own sample-size labels apply.
"""
import functools
import glob
import inspect
import os
import re
from datetime import datetime

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover -- same optional-import convention as backtest_signal_pnl.py
    load_workbook = None

LOG_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "signal_logs")
TRACKER_FILENAME = "gamma_strategy_signals_clean.xlsx"
TRACKER_SHEET = "Gamma Trade Tracking"
DAILY_SHEET = "Gamma Signals"
_DAILY_NAME = re.compile(r"^gamma_blast_\d{4}-\d{2}-\d{2}\.xlsx$")   # skips *_pre-update / *_corrupted_* archives

# status values the live engine treats as finished (see gamma_microstructure._update_alert_lifecycles)
TERMINAL = {"TARGET_2_HIT", "STOPPED_OUT", "TARGET_1_HIT_TRAILED", "EXPIRED"}
REPORT_TITLE = "F&O Sniper -- Gamma Blast Strategy Backtest"
INDEX_LABEL = "Gamma Blast Strategy"     # used only if the PDF engine on disk predates the report_title option


# ---------------------------------------------------------------------------
# small parsers
# ---------------------------------------------------------------------------
def _num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if f == f else None             # NaN -> None


def _parse_dt(v):
    """'2026-09-24 10:15:30 IST' (or a real datetime) -> naive datetime; None if blank/unparseable."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    s = str(v).strip()
    if s.upper().endswith(" IST"):
        s = s[:-4].strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def _norm_date(v):
    s = str(v or "").strip()[:10]
    try:
        datetime.strptime(s, "%Y-%m-%d")
        return s
    except ValueError:
        return None


def _ckey(r):
    c = str(r.get("contract") or "").strip().upper()
    if c:
        return c
    return "|".join(str(r.get(k) or "").strip().upper() for k in ("symbol", "expiry", "strike", "option_type"))


# ---------------------------------------------------------------------------
# reading the two workbooks
# ---------------------------------------------------------------------------
def _tracker_path():
    return os.environ.get("GAMMA_TRADE_TRACKER_XLSX") or os.path.join(LOG_DIR, TRACKER_FILENAME)


def _read_sheet(path, sheet):
    """[{header: value}] for one sheet; [] if the file/sheet is missing or unreadable (never raises)."""
    if load_workbook is None or not os.path.exists(path):
        return []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print(f"  (gamma backtest: skipping {os.path.basename(path)}: {e})")
        return []
    try:
        if sheet not in wb.sheetnames:
            return []
        rows = wb[sheet].iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            return []
        headers = [str(h).strip() if h is not None else None for h in headers]
        out = []
        for raw in rows:
            if raw is None or all(c is None for c in raw):
                continue
            out.append({h: v for h, v in zip(headers, raw) if h})
        return out
    except Exception as e:
        print(f"  (gamma backtest: could not read {os.path.basename(path)}: {e})")
        return []
    finally:
        wb.close()


def _from_daily(r):
    return {
        "alert_id": r.get("Alert ID"), "contract": r.get("Contract"), "symbol": r.get("Symbol"),
        "option_type": r.get("Option Type"), "strike": r.get("Strike"), "expiry": r.get("Expiry"), "dte": r.get("DTE"),
        "entry": r.get("Entry Premium"), "sl": r.get("Stop Loss"), "t1": r.get("Target 1"), "t2": r.get("Target 2"),
        "lot_size": r.get("Lot Size"), "status": r.get("Status") or "ACTIVE", "realized_r": r.get("Realized R"),
        "exit_price": r.get("Exit Price"), "entered": r.get("Timestamp"), "closed": r.get("Exit Time"),
        "trigger_candle": r.get("Trigger Candle"), "trailing_sl": None,
        "oi_drop": r.get("OI Drop % (Phase 1)"), "vol_exp": r.get("Volume Expansion x (Phase 3)"),
        "price_lift": r.get("Price Lift % (Phase 4)"),
        "t1_hit": None, "t2_hit": None, "sl_hit": None, "carry": False,
    }


def _from_tracker(r):
    return {
        "alert_id": r.get("Alert ID"), "contract": r.get("Contract"), "symbol": r.get("Symbol"),
        "option_type": r.get("Option Type"), "strike": r.get("Strike"), "expiry": r.get("Expiry"), "dte": None,
        "entry": r.get("Entry Price"), "sl": r.get("Stop Loss"), "t1": r.get("Target 1"), "t2": r.get("Target 2"),
        "lot_size": None, "status": r.get("Status") or "ACTIVE", "realized_r": r.get("Realized R"),
        "exit_price": None, "entered": r.get("Timestamp (IST)"), "closed": r.get("Closed At (IST)"),
        "trigger_candle": r.get("Trigger Candle"), "trailing_sl": r.get("Trailing SL"),
        "oi_drop": None, "vol_exp": None, "price_lift": None,
        "t1_hit": r.get("T1 Hit At (IST)"), "t2_hit": r.get("T2 Hit At (IST)"), "sl_hit": r.get("SL Hit At (IST)"),
        "carry": str(r.get("Carry Forward") or "").strip().upper() == "YES",   # YES only while still open overnight; the report derives "carried overnight" itself from entry/exit dates
    }


def collect_rows():
    """One merged dict per Gamma contract (see the module docstring for the merge rule)."""
    daily = []
    for f in sorted(glob.glob(os.path.join(LOG_DIR, "*", "gamma_blast_*.xlsx"))):
        if not _DAILY_NAME.match(os.path.basename(f)):
            continue
        daily += [_from_daily(r) for r in _read_sheet(f, DAILY_SHEET) if r.get("Contract")]

    by_alert, by_contract = {}, {}
    for d in daily:
        if d["alert_id"]:
            by_alert[d["alert_id"]] = d
        by_contract.setdefault(_ckey(d), d)                # files are read oldest first -> earliest trigger wins

    merged, used = [], set()
    for t in (_from_tracker(r) for r in _read_sheet(_tracker_path(), TRACKER_SHEET) if r.get("Contract")):
        d = by_alert.get(t["alert_id"]) if t["alert_id"] else None
        d = d or by_contract.get(_ckey(t))
        row = dict(t)
        if d:
            used.add(id(d))
            for k in ("lot_size", "oi_drop", "vol_exp", "price_lift", "dte", "trigger_candle"):
                if row.get(k) in (None, ""):
                    row[k] = d.get(k)
            # The tracker keeps updating a trade across days, so it wins. Only if it still says "open" while the
            # day's own file already recorded the finish do we take the finish from the daily file.
            if str(row["status"]).upper() not in TERMINAL and str(d["status"]).upper() in TERMINAL:
                for k in ("status", "realized_r", "exit_price", "closed"):
                    row[k] = d[k]
        merged.append(row)

    seen = {_ckey(r) for r in merged}
    for d in daily:
        if id(d) in used or _ckey(d) in seen:
            continue
        seen.add(_ckey(d))
        merged.append(d)
    return merged


# ---------------------------------------------------------------------------
# row -> trade (the dict shape backtest_signal_pnl.py's engine expects)
# ---------------------------------------------------------------------------
def _realized_r(row, status, entry, sl, risk):
    """(R, derived). The logged Realized R is used whenever it is there; otherwise it is recomputed from the
    engine's own rules -- but only where the answer is unambiguous. None = cannot be recovered."""
    logged = _num(row.get("realized_r"))
    if logged is not None and abs(logged) > 1e-9:                # 0.0 is not a legal finished outcome for any status
        return logged, False
    t1, t2 = _num(row.get("t1")), _num(row.get("t2"))
    if status == "STOPPED_OUT":
        return (sl - entry) / risk, True
    if status == "TARGET_1_HIT_TRAILED" and t1 is not None:
        trail = _num(row.get("trailing_sl"))
        trail = entry if trail is None else trail
        return 0.5 * (t1 - entry) / risk + 0.5 * (trail - entry) / risk, True
    if status == "TARGET_2_HIT" and t1 is not None and t2 is not None and row.get("t1_hit"):
        return 0.5 * (t1 - entry) / risk + 0.5 * (t2 - entry) / risk, True
    if status == "EXPIRED":
        px = _num(row.get("exit_price"))
        if px is not None:
            return (px - entry) / risk, True
    return None, True


def _build_trade(row, lot_for_symbol):
    """-> (trade | None, reason). reason is None for a trade, else a short key for why it was left out."""
    status = str(row.get("status") or "ACTIVE").strip().upper()
    if status not in TERMINAL:
        return None, "open"
    entry, sl = _num(row.get("entry")), _num(row.get("sl"))
    if entry is None or entry <= 0 or sl is None or sl <= 0 or sl >= entry:
        return None, "bad_levels"
    otype = str(row.get("option_type") or "").strip().upper()
    if otype not in ("CE", "PE"):
        return None, "bad_option_type"
    entry_dt = _parse_dt(row.get("entered"))
    if entry_dt is None:
        return None, "no_entry_time"

    risk = max(0.05, entry - sl)                                  # the live engine's own floor
    r, derived = _realized_r(row, status, entry, sl, risk)
    if r is None:
        return None, "r_unrecoverable"
    exit_px = max(0.0, entry + r * risk)          # NOT rounded: a blended two-leg exit is not a real tick price, and rounding it to 2 dp would shift P&L on cheap premiums

    exit_dt = _parse_dt(row.get("closed"))
    if exit_dt is None and status == "STOPPED_OUT":
        exit_dt = _parse_dt(row.get("sl_hit"))
    if exit_dt is None and status == "TARGET_2_HIT":
        exit_dt = _parse_dt(row.get("t2_hit"))
    if exit_dt is None:
        return None, "no_exit_time"
    if exit_dt < entry_dt:
        return None, "exit_before_entry"

    symbol = str(row.get("symbol") or "").strip().upper()
    lot = _num(row.get("lot_size"))
    qty = int(lot) if lot and lot > 1 else lot_for_symbol(symbol)   # a logged 1 is the restart-rebuild placeholder, not a real lot
    if not qty:
        return None, "no_lot_size"

    t1, t2 = _num(row.get("t1")), _num(row.get("t2"))
    if status == "STOPPED_OUT":
        reason = "SL Hit"
    elif status == "TARGET_2_HIT":
        reason = "Target 2 Hit"
    elif status == "TARGET_1_HIT_TRAILED":
        reason = "Target 1 Hit"
    else:
        reason = "Closed (expired)"

    pnl = round(qty * (exit_px - entry), 2)
    return {
        "symbol": symbol, "action": "BUY" if otype == "CE" else "SELL",   # same convention as the Sniper report: long call = BUY, long put = SELL
        "grade": None, "sector": None, "oi_confirmation": None,
        "pattern": row.get("trigger_candle") or "Unknown", "confidence": None,
        "entry_dt": entry_dt, "exit_dt": exit_dt,
        "entry": entry, "sl": sl, "exit_price": exit_px, "qty": qty,
        "target1": t1, "target2": t2, "target3": None,
        "pnl": pnl, "pnl_pct": round((exit_px - entry) / entry * 100, 2), "exit_reason": reason,
        "r_multiple": round(r, 3),
        "sl_hit_at": _parse_dt(row.get("sl_hit")), "target1_hit_at": _parse_dt(row.get("t1_hit")),
        "target2_hit_at": _parse_dt(row.get("t2_hit")), "target3_hit_at": None,
        "india_vix_at_signal": None, "expiry_date": _norm_date(row.get("expiry")),
        # Gamma-only context (the engine ignores extra keys)
        "contract": row.get("contract"), "option_type": otype, "carried_overnight": exit_dt.date() > entry_dt.date(),
        "dte": row.get("dte"), "oi_drop_pct": _num(row.get("oi_drop")), "volume_expansion_x": _num(row.get("vol_exp")),
        "price_lift_pct": _num(row.get("price_lift")), "r_was_derived": derived,
    }, None


def load_gamma_trades(rows=None):
    """
    -> (trades, excluded, info). trades are finished Gamma trades, sorted by exit time. excluded counts every
    logged contract that did not become a trade (still-open ones included, same definition the Sniper report
    uses); info breaks that down by reason so nothing is silently lost.
    """
    rows = collect_rows() if rows is None else rows

    def lot_for_symbol(sym):
        try:
            from .lot_size_resolver import get_lot_size
            return get_lot_size(sym) if sym else None
        except Exception:
            return None

    trades, reasons = [], {}
    for row in rows:
        t, why = _build_trade(row, lot_for_symbol)
        if t is not None:
            trades.append(t)
        else:
            reasons[why] = reasons.get(why, 0) + 1
    trades.sort(key=lambda t: t["exit_dt"])
    info = {
        "logged_contracts": len(rows), "resolved_trades": len(trades), "still_open": reasons.get("open", 0),
        "excluded_other": {k: v for k, v in reasons.items() if k != "open"},
        "r_recomputed": sum(1 for t in trades if t["r_was_derived"]),
    }
    return trades, len(rows) - len(trades), info


# ---------------------------------------------------------------------------
# report / API helpers (called from daily_backtest.py and gamma_backtest_views.py)
# ---------------------------------------------------------------------------
def _write_pdf(trades, metrics, excluded, final_name):
    from .backtest_signal_pnl import write_pdf_report
    from .daily_backtest import _run_and_rename
    kwargs = dict(is_index=True, index_name=INDEX_LABEL, excluded_count=excluded)   # is_index: skips the Grade/Sector/OI sections Gamma has no data for
    if "report_title" in inspect.signature(write_pdf_report).parameters:
        kwargs["report_title"] = REPORT_TITLE
    return _run_and_rename(functools.partial(write_pdf_report, **kwargs), trades, metrics, final_name)


def run_gamma_cycle(end_str):
    """The Gamma step of the daily cycle. Always returns the same keys; every value is None when there is
    nothing resolved yet (expected early on -- not an error). Summary / curve / trades are exposed BEFORE the PDF
    is attempted, so a PDF-library problem only costs the download link, never the numbers."""
    from .backtest_signal_pnl import compute_capital_base, compute_metrics, compute_equity_curve
    from .daily_backtest import _summarize, _serialize_equity_curve, _serialize_trades
    out = {"pdf": None, "summary": None, "equity_curve": None, "recent_trades": None, "pdf_error": None, "info": None}
    trades, excluded, info = load_gamma_trades()
    out["info"] = info
    if not trades:
        return out
    capital_base = compute_capital_base(trades)               # premium buying: real peak (lot x premium) committed, like the Sniper report
    metrics = compute_metrics(trades, capital_base)
    if not metrics:
        return out
    out["summary"] = _summarize(metrics)
    out["equity_curve"] = _serialize_equity_curve(compute_equity_curve(trades, capital_base))
    out["recent_trades"] = _serialize_trades(trades)
    try:
        out["pdf"] = _write_pdf(trades, metrics, excluded, f"gamma_blast_{end_str}.pdf")
    except Exception as e:
        out["pdf_error"] = str(e)
    return out


def run_gamma_range_preview(start_date, end_date):
    """JSON preview for the date-range picker: {'summary','equity_curve','recent_trades'}, values None if empty."""
    from .backtest_signal_pnl import compute_capital_base, compute_metrics, compute_equity_curve
    from .daily_backtest import _summarize, _serialize_equity_curve, _serialize_trades, _filter_trades_by_range
    result = {"summary": None, "equity_curve": None, "recent_trades": None}
    trades, _, _ = load_gamma_trades()
    trades = _filter_trades_by_range(trades, start_date, end_date)
    if trades:
        capital_base = compute_capital_base(trades)
        metrics = compute_metrics(trades, capital_base)
        if metrics:
            result["summary"] = _summarize(metrics)
            result["equity_curve"] = _serialize_equity_curve(compute_equity_curve(trades, capital_base))
            result["recent_trades"] = _serialize_trades(trades)
    return result


def run_gamma_range_report(start_str, end_str):
    """Downloadable Gamma PDF for [start, end] (by exit date). ValueError on a bad date, None if nothing resolved."""
    try:
        start_date = datetime.strptime(start_str, "%Y-%m-%d").date()
        end_date = datetime.strptime(end_str, "%Y-%m-%d").date()
    except (ValueError, TypeError):
        raise ValueError("start/end must be YYYY-MM-DD")
    from .backtest_signal_pnl import compute_capital_base, compute_metrics
    from .daily_backtest import _filter_trades_by_range
    trades, excluded, _ = load_gamma_trades()
    trades = _filter_trades_by_range(trades, start_date, end_date)
    if not trades:
        return None
    metrics = compute_metrics(trades, compute_capital_base(trades))
    if not metrics:
        return None
    return _write_pdf(trades, metrics, excluded, f"gamma_blast_{start_str}_to_{end_str}.pdf")


if __name__ == "__main__":
    import sys
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from screener.backtest_signal_pnl import compute_capital_base, compute_metrics, print_summary

    if load_workbook is None:
        print("openpyxl not installed -- pip install openpyxl")
    else:
        all_trades, skipped, details = load_gamma_trades()
        print(f"Gamma trade log: {details}")
        if not all_trades:
            print("No finished Gamma trades yet -- not an error, just nothing resolved so far.")
        else:
            m = compute_metrics(all_trades, compute_capital_base(all_trades))
            print_summary(m, skipped)
            try:
                print(f"\nFull PDF report written to: {run_gamma_cycle(datetime.now().strftime('%Y-%m-%d'))['pdf']}")
            except ImportError:
                print("\nPDF generation needs matplotlib and reportlab -- pip install matplotlib reportlab")
