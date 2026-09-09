"""
Backfill signal outcomes from real Fyers 1-minute candles.

Signal outcomes are positional, not forced to EOD. A signal remains eligible
for SL/Target resolution on later trading days until the requested end date
or the option contract expiry, whichever comes first.

The consolidated workbook contains two sheets:
  Results  - genuine SL/Target hits only.
  No Result - no verified SL/Target hit found by the last date checked.

An EOD price move is never converted into "Closed up/down" because that is
not a genuine SL/Target outcome for a positional signal.

Target progression is preserved. For example, if Target 1 is reached on one
session and the original SL is reached later, the final Result is still
"SL Hit", but Trade Progression records "Target 1 Hit -> SL Hit". This makes
it possible to distinguish a trade that never reached a target from one that
reached targets before subsequently reversing into its stop.
"""
import os
from datetime import datetime

from django.core.management.base import BaseCommand
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from screener.excel_logger import COLUMNS, LOG_DIR
from screener.fyers_client import get_history


_HISTORY_CACHE = {}


def _find_log_path(date_str):
    nested = os.path.join(LOG_DIR, date_str, f"signals_{date_str}.xlsx")
    if os.path.exists(nested):
        return nested
    flat = os.path.join(LOG_DIR, f"signals_{date_str}.xlsx")
    if os.path.exists(flat):
        return flat
    return None


def _parse_datetime(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value), fmt)
        except Exception:
            pass
    return None


def _fetch_candles(option_symbol, start_date, end_date):
    """Fetch one multi-day positional range; cache identical requests."""
    key = (option_symbol, str(start_date), str(end_date))
    if key in _HISTORY_CACHE:
        return _HISTORY_CACHE[key]

    resp = get_history(
        option_symbol,
        resolution="1",
        range_from=str(start_date),
        range_to=str(end_date),
    )
    if not resp or resp.get("s") != "ok":
        value = (None, resp.get("message") if resp else "no response")
    else:
        value = (resp.get("candles", []), None)
    _HISTORY_CACHE[key] = value
    return value


def _replay(candles, entry_dt, sl, t1, t2, t3):
    """Replay positional SL/target crossings chronologically."""
    result = {"sl_hit_at": None, "targets_hit": {}, "ambiguous": []}
    targets = ((1, t1), (2, t2), (3, t3))

    for candle in sorted(candles, key=lambda c: c[0]):
        if len(candle) < 5:
            continue
        ts, _open, high, low, _close = candle[:5]
        candle_dt = datetime.fromtimestamp(ts)
        if candle_dt < entry_dt:
            continue

        sl_touched = sl is not None and low <= sl
        newly_touched = [
            (n, level)
            for n, level in targets
            if n not in result["targets_hit"] and level is not None and high >= level
        ]

        if sl_touched and newly_touched:
            result["ambiguous"].append(candle_dt.strftime("%Y-%m-%d %H:%M:%S"))
            return result

        if sl_touched:
            result["sl_hit_at"] = candle_dt.strftime("%Y-%m-%d %H:%M:%S")
            return result

        if newly_touched:
            hit_at = candle_dt.strftime("%Y-%m-%d %H:%M:%S")
            for n, _level in newly_touched:
                result["targets_hit"][n] = hit_at
            if 3 in result["targets_hit"]:
                return result

    return result


def _progression_from_hits(sl_hit_at, targets_hit):
    """Return a chronological human-readable progression from verified hits."""
    events = []
    for n, hit_at in targets_hit.items():
        if hit_at:
            events.append((str(hit_at), f"Target {n} Hit"))
    if sl_hit_at:
        events.append((str(sl_hit_at), "SL Hit"))
    events.sort(key=lambda item: item[0])
    if not events:
        return ""
    return " -> ".join(label for _ts, label in events)


def _max_target(targets_hit):
    return max(targets_hit) if targets_hit else 0


def _pnl_pct(outcome, entry, sl, t1, t2, t3):
    if not outcome or entry in (None, 0):
        return None
    if outcome == "SL Hit" and sl is not None:
        return round((sl - entry) / entry * 100, 2)
    if outcome.startswith("Target"):
        try:
            n = int(outcome.split()[1])
        except Exception:
            return None
        level = {1: t1, 2: t2, 3: t3}.get(n)
        return round((level - entry) / entry * 100, 2) if level is not None else None
    return None


def _report_headers(ws):
    headers = [
        "Date", "Symbol", "Action", "Grade", "Strike",
        "Entry (Premium)", "SL", "Target 1", "Target 2", "Target 3",
        "Option Symbol", "Result", "Trade Progression", "Max Target Reached",
        "Hit At", "P&L %", "Checked Through",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")


class Command(BaseCommand):
    help = "Resolve positional SL/Target outcomes through later sessions/expiry and build a two-sheet report."

    def add_arguments(self, parser):
        parser.add_argument("--start", required=True, help="YYYY-MM-DD")
        parser.add_argument("--end", required=True, help="YYYY-MM-DD")
        parser.add_argument("--dry-run", action="store_true", help="Preview without writing signal logs/report.")

    def handle(self, *args, **options):
        global _HISTORY_CACHE
        _HISTORY_CACHE = {}

        start = datetime.strptime(options["start"], "%Y-%m-%d").date()
        requested_end = datetime.strptime(options["end"], "%Y-%m-%d").date()
        dry_run = options["dry_run"]
        results = []
        pending = []

        d = start
        while d <= requested_end:
            date_str = d.strftime("%Y-%m-%d")
            path = _find_log_path(date_str)
            if not path:
                d = d.fromordinal(d.toordinal() + 1)
                continue

            self.stdout.write(f"\n--- {date_str} ({os.path.basename(path)}) ---")
            wb = load_workbook(path)
            ws = wb["Signals"]
            headers = [c.value for c in ws[1]]
            if headers != COLUMNS:
                self.stdout.write(self.style.WARNING("  Column layout doesn't match current schema -- skipping."))
                d = d.fromordinal(d.toordinal() + 1)
                continue

            col = {name: i + 1 for i, name in enumerate(COLUMNS)}
            changed = 0
            checked = 0

            for row_num in range(2, ws.max_row + 1):
                symbol = ws.cell(row_num, col["Symbol"]).value
                action = ws.cell(row_num, col["Action"]).value
                grade = ws.cell(row_num, col["Grade"]).value
                outcome = ws.cell(row_num, col["Outcome"]).value
                opt_symbol = ws.cell(row_num, col["Option Symbol"]).value
                entry = ws.cell(row_num, col["Entry (Premium)"]).value
                strike = ws.cell(row_num, col["Strike"]).value
                sl = ws.cell(row_num, col["SL"]).value
                t1 = ws.cell(row_num, col["Target 1"]).value
                t2 = ws.cell(row_num, col["Target 2"]).value
                t3 = ws.cell(row_num, col["Target 3"]).value
                ts_raw = ws.cell(row_num, col["Timestamp"]).value
                expiry_raw = ws.cell(row_num, col["Expiry Date"]).value

                if not symbol or not opt_symbol or entry in (None, 0) or None in (sl, t1, t2, t3) or not ts_raw:
                    continue

                entry_dt = _parse_datetime(ts_raw)
                if not entry_dt:
                    continue

                existing_sl_hit = ws.cell(row_num, col["SL Hit At"]).value
                existing_targets = {
                    n: ws.cell(row_num, col[f"Target {n} Hit At"]).value
                    for n in (1, 2, 3)
                }
                genuine_hit = bool(existing_sl_hit or any(existing_targets.values()))

                # Remove only old synthetic EOD outcomes. Never erase a genuine
                # SL/Target hit recorded by the live logger.
                if outcome and not genuine_hit and (
                    str(outcome).startswith("Closed ")
                    or str(outcome).startswith("No trade data")
                    or str(outcome).startswith("Unresolved")
                ):
                    if not dry_run:
                        ws.cell(row_num, col["Outcome"]).value = ""
                    outcome = ""
                    changed += 1

                if genuine_hit:
                    progression = _progression_from_hits(existing_sl_hit, existing_targets)
                    max_target = _max_target({n: v for n, v in existing_targets.items() if v})
                    hit_at = existing_sl_hit or existing_targets.get(3) or existing_targets.get(2) or existing_targets.get(1)
                    results.append({
                        "date": date_str, "symbol": symbol, "action": action, "grade": grade,
                        "strike": strike, "entry": entry, "sl": sl, "t1": t1, "t2": t2, "t3": t3,
                        "option_symbol": opt_symbol, "result": outcome or "Resolved",
                        "progression": progression, "max_target": max_target,
                        "hit_at": hit_at, "pnl_pct": _pnl_pct(outcome, entry, sl, t1, t2, t3),
                        "checked_through": date_str,
                    })
                    continue

                expiry_date = _parse_datetime(expiry_raw)
                terminal_date = requested_end
                if expiry_date and expiry_date.date() < terminal_date:
                    terminal_date = expiry_date.date()
                if terminal_date < entry_dt.date():
                    terminal_date = entry_dt.date()

                checked += 1
                candles, err = _fetch_candles(opt_symbol, entry_dt.date(), terminal_date)
                if candles is None:
                    pending.append({
                        "date": date_str, "symbol": symbol, "action": action, "grade": grade,
                        "strike": strike, "entry": entry, "sl": sl, "t1": t1, "t2": t2, "t3": t3,
                        "option_symbol": opt_symbol,
                        "status": f"No verified result — historical data unavailable ({err or 'unknown'})",
                        "checked_through": str(terminal_date),
                    })
                    continue

                if not candles:
                    pending.append({
                        "date": date_str, "symbol": symbol, "action": action, "grade": grade,
                        "strike": strike, "entry": entry, "sl": sl, "t1": t1, "t2": t2, "t3": t3,
                        "option_symbol": opt_symbol,
                        "status": "No verified SL/Target result — no historical candles returned",
                        "checked_through": str(terminal_date),
                    })
                    continue

                replay = _replay(candles, entry_dt, sl, t1, t2, t3)
                if replay["ambiguous"]:
                    pending.append({
                        "date": date_str, "symbol": symbol, "action": action, "grade": grade,
                        "strike": strike, "entry": entry, "sl": sl, "t1": t1, "t2": t2, "t3": t3,
                        "option_symbol": opt_symbol,
                        "status": f"AMBIGUOUS — SL and target touched in same 1-min candle at {replay['ambiguous'][0]}",
                        "checked_through": str(terminal_date),
                    })
                    continue

                targets_hit = replay["targets_hit"]
                max_target = _max_target(targets_hit)
                progression = _progression_from_hits(replay["sl_hit_at"], targets_hit)

                if replay["sl_hit_at"]:
                    new_outcome = "SL Hit"
                    hit_at = replay["sl_hit_at"]
                    if not dry_run:
                        for n, hit in targets_hit.items():
                            ws.cell(row_num, col[f"Target {n} Hit At"]).value = hit
                        ws.cell(row_num, col["SL Hit At"]).value = hit_at
                        ws.cell(row_num, col["Outcome"]).value = new_outcome
                    changed += 1
                    results.append({
                        "date": date_str, "symbol": symbol, "action": action, "grade": grade,
                        "strike": strike, "entry": entry, "sl": sl, "t1": t1, "t2": t2, "t3": t3,
                        "option_symbol": opt_symbol, "result": new_outcome, "progression": progression,
                        "max_target": max_target, "hit_at": hit_at,
                        "pnl_pct": _pnl_pct(new_outcome, entry, sl, t1, t2, t3),
                        "checked_through": str(terminal_date),
                    })
                elif max_target:
                    new_outcome = f"Target {max_target} Hit"
                    hit_at = targets_hit[max_target]
                    if not dry_run:
                        for n, hit in targets_hit.items():
                            ws.cell(row_num, col[f"Target {n} Hit At"]).value = hit
                        ws.cell(row_num, col["Outcome"]).value = new_outcome
                    changed += 1
                    results.append({
                        "date": date_str, "symbol": symbol, "action": action, "grade": grade,
                        "strike": strike, "entry": entry, "sl": sl, "t1": t1, "t2": t2, "t3": t3,
                        "option_symbol": opt_symbol, "result": new_outcome, "progression": progression,
                        "max_target": max_target, "hit_at": hit_at,
                        "pnl_pct": _pnl_pct(new_outcome, entry, sl, t1, t2, t3),
                        "checked_through": str(terminal_date),
                    })
                else:
                    expiry_note = " — contract expiry reached" if expiry_date and terminal_date == expiry_date.date() else ""
                    pending.append({
                        "date": date_str, "symbol": symbol, "action": action, "grade": grade,
                        "strike": strike, "entry": entry, "sl": sl, "t1": t1, "t2": t2, "t3": t3,
                        "option_symbol": opt_symbol,
                        "status": f"No SL/Target hit through {terminal_date}{expiry_note}",
                        "checked_through": str(terminal_date),
                    })

            if changed and not dry_run:
                wb.save(path)
            self.stdout.write(self.style.SUCCESS(f"  {changed} row(s) updated; {checked} positional rows checked."))
            d = d.fromordinal(d.toordinal() + 1)

        self._write_report(results, pending, options["start"], options["end"], dry_run)

    def _write_report(self, results, pending, start_str, end_str, dry_run):
        wb = Workbook()
        ws_results = wb.active
        ws_results.title = "Results"
        _report_headers(ws_results)
        for r in results:
            ws_results.append([
                r["date"], r["symbol"], r["action"], r["grade"], r["strike"],
                r["entry"], r["sl"], r["t1"], r["t2"], r["t3"], r["option_symbol"],
                r["result"], r["progression"], r["max_target"], r["hit_at"],
                r["pnl_pct"], r["checked_through"],
            ])

        ws_pending = wb.create_sheet("No Result")
        _report_headers(ws_pending)
        for r in pending:
            ws_pending.append([
                r["date"], r["symbol"], r["action"], r["grade"], r["strike"],
                r["entry"], r["sl"], r["t1"], r["t2"], r["t3"], r["option_symbol"],
                r["status"], "", "", "", "", r["checked_through"],
            ])

        for sheet in (ws_results, ws_pending):
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cells in sheet.columns:
                max_len = max(len(str(c.value or "")) for c in cells)
                sheet.column_dimensions[cells[0].column_letter].width = min(max(max_len + 2, 10), 45)

        out_dir = os.path.join(LOG_DIR, "backfill_reports")
        os.makedirs(out_dir, exist_ok=True)
        path = os.path.join(out_dir, f"complete_report_{start_str}_to_{end_str}.xlsx")
        if not dry_run:
            wb.save(path)

        self.stdout.write(f"\n{'=' * 60}")
        self.stdout.write(f"  Results: {len(results)} genuine SL/Target outcomes")
        self.stdout.write(f"  No Result: {len(pending)} positions with no verified SL/Target outcome")
        self.stdout.write(f"  Saved: {path}")
        self.stdout.write(f"{'=' * 60}")

        if dry_run:
            return

        try:
            from trading.telegram_bot import TelegramBot
            bot = TelegramBot()
            caption = (
                f"📊 <b>F&O Radar — Positional Outcome Report</b>\n"
                f"{os.path.basename(path)}\n"
                f"Results: {len(results)} | No Result: {len(pending)}"
            )
            result = bot.send_document(path, caption=caption)
            if result and result.get("ok"):
                self.stdout.write(self.style.SUCCESS("Sent to Telegram."))
            else:
                self.stdout.write(self.style.WARNING(f"Report saved but Telegram send didn't confirm success: {result}"))
        except Exception as e:
            self.stdout.write(self.style.WARNING(f"Report saved but Telegram send failed: {e}"))
