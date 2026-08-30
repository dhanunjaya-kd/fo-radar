"""
screener/backtest_index_positional.py

Positional NIFTY/BANKNIFTY backtest -- treats a genuine Bias flip
(from the "Flips" sheet index_tracker.py already logs) as a real
futures position, held across multiple days until the opposite flip
happens (or a max-hold safety cap), using only data this project
already has: Bias, Support/Resistance, Fut price, and the overnight
gap computed from Spot history. Deliberately excludes crypto and
"global cues" -- neither is tracked anywhere in this project; adding
either would be new infrastructure, not this backtest.

Reuses the ENTIRE metrics/PDF engine from backtest_signal_pnl.py
(compute_metrics, compute_capital_base, generate_key_findings,
write_pdf_report and everything under it) rather than rebuilding it --
that engine is already thoroughly tested; only the trade-GENERATION
logic here is new. A positional trade is mapped into the exact same
trade dict shape (entry_dt/exit_dt/pnl/...) the stock backtest uses,
so everything downstream (Sharpe/Sortino/drawdown/segment breakdown/
PDF layout) works unmodified.

SIZING: unlike the stock backtest's Rs 50,000/trade (an existing
convention pulled from views.py), there's no equivalent existing
precedent for index futures anywhere in this codebase -- this is a
new assumption, not an established one. Uses 1 lot per trade, P&L
shown in both raw index points (the real, assumption-free number) and
an estimated rupee figure using a clearly-labelled, CONFIGURABLE lot
size (NIFTY/BANKNIFTY lot sizes change periodically per NSE circulars
-- verify the current one before trusting the rupee figure).

Run as a standalone script:
    cd backend
    venv\\Scripts\\activate
    python -m screener.backtest_index_positional
"""
import os
import glob
import re
from datetime import datetime, timedelta

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

LOG_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "signal_logs")

# NSE lot sizes drift over time (periodic NSE circulars) -- these are
# reasonable recent-ish defaults, NOT guaranteed current. The backtest
# still works correctly in POINTS regardless of whether this is stale;
# only the rupee conversion depends on it being right.
DEFAULT_LOT_SIZE = {"NIFTY": 75, "BANKNIFTY": 35}

MAX_HOLD_DAYS = 10  # safety cap -- if Bias never flips back, force-close rather than hold indefinitely


def _direction(bias):
    """Same direction-only reading backtest_index_bias.py's own
    direction() already uses -- Bullish/Bullish (Strong) -> 'up',
    Bearish/Bearish (Strong) -> 'down', Neutral/blank -> None."""
    if not bias:
        return None
    if bias.startswith("Bullish"):
        return "up"
    if bias.startswith("Bearish"):
        return "down"
    return None


def _time_to_minutes(t):
    """
    Normalizes a Snapshots/Flips 'Time' cell value into minutes-since-
    midnight, for the staleness check in _snapshot_at_or_before() below.
    Handles the normal openpyxl case (a real datetime.time or datetime
    object) directly via .hour/.minute, and falls back to parsing a
    plain "HH:MM..." string for anything else. Returns None if
    genuinely unparseable -- callers must treat that as "can't tell,
    don't block a trade over a format surprise," never raise.
    """
    if t is None:
        return None
    if hasattr(t, "hour") and hasattr(t, "minute"):
        return t.hour * 60 + t.minute
    try:
        parts = str(t).split(":")
        return int(parts[0]) * 60 + int(parts[1])
    except (ValueError, IndexError):
        return None


def list_index_tracker_dates(index_name):
    """Every date with a real index_tracker log for this index, oldest first
    (this backtest needs to walk forward chronologically, unlike the stock
    backtest which sorts by exit time after the fact)."""
    dates = []
    pattern = re.compile(rf'index_tracker_{index_name}_(\d{{4}}-\d{{2}}-\d{{2}})\.xlsx$')
    for path in glob.glob(os.path.join(LOG_DIR, "*", f"index_tracker_{index_name}_*.xlsx")):
        m = pattern.search(os.path.basename(path))
        if m:
            dates.append(m.group(1))
    return sorted(dates)


def load_flips(index_name, date_str):
    """Every row from that day's Flips sheet, oldest first. Returns []
    if the file or sheet doesn't exist (a day with zero flips, or an
    older file from before the Flips sheet existed)."""
    path = os.path.join(LOG_DIR, date_str, f"index_tracker_{index_name}_{date_str}.xlsx")
    if not os.path.exists(path):
        return []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if "Flips" not in wb.sheetnames:
            return []
        ws = wb["Flips"]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        rows = []
        for raw in ws.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, raw))
            if row.get("Date") and row.get("Time"):
                rows.append(row)
        return rows
    except Exception as e:
        print(f"  (skipping Flips in {os.path.basename(path)}: {e})")
        return []


def load_snapshots(index_name, date_str):
    """Every row from that day's Snapshots sheet, oldest first (real
    file order is newest-first for display purposes elsewhere in this
    project -- reversed here since this backtest needs chronological
    walk-forward)."""
    path = os.path.join(LOG_DIR, date_str, f"index_tracker_{index_name}_{date_str}.xlsx")
    if not os.path.exists(path):
        return []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Snapshots"]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        rows = []
        for raw in ws.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, raw))
            if row.get("Time"):
                rows.append(row)
        rows.reverse()  # file is newest-first; this backtest needs oldest-first
        return rows
    except Exception as e:
        print(f"  (skipping Snapshots in {os.path.basename(path)}: {e})")
        return []


def compute_gap_pct(all_snapshots_by_date, date_str):
    """Overnight gap for this date: (this day's FIRST Spot reading -
    previous trading day's LAST Spot reading) / previous close * 100.
    Returns None if there's no real previous day's data to compare
    against (the very first logged day) -- never fabricates a gap from
    nothing."""
    dates = sorted(all_snapshots_by_date.keys())
    idx = dates.index(date_str) if date_str in dates else None
    if idx is None or idx == 0:
        return None
    prev_date = dates[idx - 1]
    today_rows = all_snapshots_by_date.get(date_str, [])
    prev_rows = all_snapshots_by_date.get(prev_date, [])
    if not today_rows or not prev_rows:
        return None
    today_open = today_rows[0].get("Spot")
    prev_close = prev_rows[-1].get("Spot")
    if not today_open or not prev_close:
        return None
    return round((today_open - prev_close) / prev_close * 100, 2)


def generate_positional_trades(index_name, lot_size=None):
    """
    Walks every real Flip event chronologically and turns genuine
    directional flips into positional trades:
      - ENTRY: a flip into Bullish-family (long) or Bearish-family
        (short), confirmed by price being on the right side of the
        already-logged Support/Resistance for that moment (long needs
        spot > Support -- not already breaking the floor it's supposed
        to be finding; short needs spot < Resistance, the mirror case).
        A flip that fails this confirmation is skipped, not force-
        entered -- same "don't fabricate a trade the data doesn't
        support" rule as every other backtest here.
      - EXIT: the next OPPOSITE-direction flip (natural, signal-driven
        exit), or MAX_HOLD_DAYS with no reversal (force-closed at the
        last available Fut price -- a real safety cap, not an
        indefinite hold).
      - P&L computed in FUTURES POINTS (not Spot -- Fut is the real
        tradeable instrument), qty = 1 lot, using lot_size purely for
        the rupee estimate alongside the points figure.

    Returns a list of trade dicts in the SAME shape
    backtest_signal_pnl.py's trades use (entry_dt/exit_dt/pnl/...) so
    the rest of that file's engine works unmodified. Also returns
    excluded_count (flips that didn't pass S/R confirmation).
    """
    # Aug 27 2026: resolve the REAL, currently-live lot size from
    # Fyers' own symbol master (lot_size_resolver.py) instead of
    # trusting DEFAULT_LOT_SIZE unconditionally -- that table was
    # already confirmed stale the moment this fix was built (a live
    # check showed BANKNIFTY's real lot size is 30, not the 35 this
    # file had hardcoded), exactly the kind of drift this function's
    # own docstring already warned about. DEFAULT_LOT_SIZE is now only
    # the LAST-RESORT fallback if the live resolver genuinely can't
    # reach Fyers (network down, or this specific run predates any
    # successful fetch) -- not the primary source anymore.
    if lot_size is None:
        try:
            from .lot_size_resolver import get_lot_size
            lot_size = get_lot_size(index_name)
        except Exception as e:
            print(f"[PositionalBacktest] Live lot-size resolve failed for {index_name}: {e}")
        if lot_size is None:
            lot_size = DEFAULT_LOT_SIZE.get(index_name, 50)
            print(f"[PositionalBacktest] Using stale fallback lot size for {index_name}: {lot_size} -- live resolve unavailable")
    dates = list_index_tracker_dates(index_name)
    if not dates:
        return [], 0

    all_snapshots_by_date = {d: load_snapshots(index_name, d) for d in dates}

    all_flips = []
    for d in dates:
        for f in load_flips(index_name, d):
            all_flips.append(f)
    all_flips.sort(key=lambda f: (str(f["Date"]), str(f["Time"])))

    def _snapshot_at_or_before(date_str, time_str, max_staleness_minutes=10):
        """The real logged snapshot closest to (at or just before) a
        given flip's timestamp, for reading the concurrent Fut price
        and Support/Resistance -- the Flips sheet's own 'Price' field
        is Spot, not Fut, so this is needed for correct P&L pricing.

        Aug 30 2026: also rejects a match that's more than
        max_staleness_minutes OLDER than the flip itself, returning
        None (same as "no snapshot at all," which the caller already
        excludes) rather than silently using it. Real bug this fixes,
        found by tracing actual reported output, not guessed: a real
        NIFTY backtest showed 4 consecutive flips across a single day
        (09:29/09:46/12:19/12:46) all pricing to the EXACT same Fut
        value, producing four artificial 0-P&L trades in a row. With a
        normal ~60s snapshot cadence, flips 17 minutes to 2.5 hours
        apart should never land on the same snapshot row -- this only
        happens if the Snapshots sheet has a real logging gap covering
        that whole stretch (this project already has one confirmed,
        still-unexplained gap pattern around the CAS window most days;
        this suggests gaps aren't limited to just that one window). 10
        minutes is a generous multiple of the normal ~60s cadence --
        wide enough to absorb ordinary jitter, tight enough to catch a
        real multi-hour gap. A rejected flip counts toward `excluded`
        via the same path a genuinely-missing snapshot already does --
        this is not a new failure mode, just a stale match no longer
        being mistaken for a fresh one.
        """
        rows = all_snapshots_by_date.get(date_str, [])
        candidate = None
        for r in rows:
            if str(r.get("Time", "")) <= str(time_str):
                candidate = r
            else:
                break
        if candidate is None:
            candidate = rows[0] if rows else None
        if candidate is None:
            return None
        flip_min = _time_to_minutes(time_str)
        snap_min = _time_to_minutes(candidate.get("Time"))
        if flip_min is not None and snap_min is not None and (flip_min - snap_min) > max_staleness_minutes:
            return None
        return candidate

    def _last_snapshot_on_or_before(date_str):
        """Latest available snapshot at or before this date -- used for
        the max-hold force-close, walking backward if a date has no
        logged data at all (a real gap in logging, not assumed)."""
        idx = dates.index(date_str) if date_str in dates else len(dates) - 1
        for d in dates[idx::-1]:
            rows = all_snapshots_by_date.get(d, [])
            if rows:
                return d, rows[-1]
        return None, None

    trades = []
    excluded = 0
    open_position = None  # {'direction', 'entry_date', 'entry_time', 'entry_fut'}

    def _close(exit_date, exit_time, exit_fut, reason):
        nonlocal open_position
        trades.append(_build_trade(index_name, open_position, exit_date, exit_time, exit_fut, reason, lot_size, all_snapshots_by_date))
        open_position = None

    flips_by_date = {}
    for f in all_flips:
        flips_by_date.setdefault(str(f["Date"]), []).append(f)

    # Walk DAY BY DAY (not just flip-event by flip-event) -- a position
    # held so long it eventually blows past MAX_HOLD_DAYS needs catching
    # on the day that happens, not only if it's STILL open once every
    # flip has been processed. Checking only at the very end (the first
    # version of this) let a position ride for 42 real days before a
    # genuine opposite flip closed it, silently ignoring the cap the
    # whole time it was in the middle of the sequence -- caught by
    # actually running this against real synthetic data, not assumed.
    for date_str in dates:
        if open_position is not None:
            entry_idx = dates.index(open_position["entry_date"])
            days_held = dates.index(date_str) - entry_idx
            if days_held >= MAX_HOLD_DAYS:
                close_date, close_row = _last_snapshot_on_or_before(date_str)
                if close_row and close_row.get("Fut") is not None:
                    _close(close_date, str(close_row.get("Time", "15:30:00")), close_row["Fut"], "Max hold reached")

        for flip in flips_by_date.get(date_str, []):
            to_dir = _direction(flip.get("To Bias"))
            time_str = str(flip["Time"])
            snap = _snapshot_at_or_before(date_str, time_str)
            if snap is None:
                excluded += 1
                continue

            if open_position and to_dir and to_dir != open_position["direction"]:
                exit_fut = snap.get("Fut")
                if exit_fut is not None:
                    _close(date_str, time_str, exit_fut, "Bias flipped opposite")

            if to_dir is None:
                continue

            if open_position is not None:
                continue

            support, resistance, entry_fut = snap.get("Support"), snap.get("Resistance"), snap.get("Fut")
            spot = snap.get("Spot")
            if entry_fut is None or spot is None:
                excluded += 1
                continue

            confirmed = (to_dir == "up" and support is not None and spot > support) or \
                        (to_dir == "down" and resistance is not None and spot < resistance)
            if not confirmed:
                excluded += 1
                continue

            open_position = {"direction": to_dir, "entry_date": date_str, "entry_time": time_str, "entry_fut": entry_fut}

    # Anything still open once real logged history simply ends (not a
    # max-hold breach, just no more data yet).
    if open_position:
        close_date, close_row = _last_snapshot_on_or_before(dates[-1])
        if close_row and close_row.get("Fut") is not None:
            _close(close_date, str(close_row.get("Time", "15:30:00")), close_row["Fut"], "End of logged history")

    return trades, excluded


def _build_trade(index_name, open_position, exit_date, exit_time, exit_fut, reason, lot_size, all_snapshots_by_date):
    """Turns one open_position + its resolved exit into the standard
    trade dict shape backtest_signal_pnl.py's engine expects."""
    direction = open_position["direction"]
    entry_fut = open_position["entry_fut"]
    points = round((exit_fut - entry_fut) if direction == "up" else (entry_fut - exit_fut), 2)
    pnl_rs = round(points * lot_size, 2)
    entry_dt = datetime.strptime(f"{open_position['entry_date']} {open_position['entry_time']}", "%Y-%m-%d %H:%M:%S")
    exit_dt = datetime.strptime(f"{exit_date} {exit_time}", "%Y-%m-%d %H:%M:%S")
    gap_pct = compute_gap_pct(all_snapshots_by_date, open_position["entry_date"])
    return {
        "symbol": index_name, "action": "BUY" if direction == "up" else "SELL", "grade": None,
        "sector": None, "oi_confirmation": None,
        "pattern": f"Gap {'up' if (gap_pct or 0) > 0 else 'down' if (gap_pct or 0) < 0 else 'flat'}" if gap_pct is not None else "Unknown",
        "entry_dt": entry_dt, "exit_dt": exit_dt,
        "entry": entry_fut, "exit_price": exit_fut, "qty": lot_size,
        "pnl": pnl_rs, "pnl_pct": round(points / entry_fut * 100, 2) if entry_fut else None,
        "exit_reason": reason, "points": points, "gap_pct_at_entry": gap_pct,
    }


# Margin (not full notional) is what real capital gets tied up per
# lot -- unlike the stock backtest's Rs 50,000, which comes from an
# existing convention already in views.py, there's no equivalent
# precedent anywhere in this codebase for index futures margin. These
# are illustrative SPAN+exposure ballpark figures, NOT a precise
# broker calculation (real margin moves with volatility) -- verify
# against your actual broker's margin requirement before trusting the
# CAGR/Sharpe/Calmar figures that get computed against this base. The
# points-based P&L above doesn't depend on this being exactly right;
# only the capital-relative ratios do.
DEFAULT_MARGIN_PER_LOT = {"NIFTY": 200000, "BANKNIFTY": 220000}


if __name__ == "__main__":
    import sys
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from screener.backtest_signal_pnl import (compute_capital_base, compute_metrics,
                                               generate_key_findings, print_summary, write_pdf_report)

    if load_workbook is None:
        print("openpyxl not installed -- pip install openpyxl")
    else:
        for index_name in ("NIFTY", "BANKNIFTY"):
            print(f"\n{'=' * 60}\n  {index_name} Positional Backtest\n{'=' * 60}")
            trades, excluded = generate_positional_trades(index_name)
            if not trades:
                print(f"No positional trades generated yet for {index_name} -- either no real flips logged, "
                      f"or none passed Support/Resistance confirmation. Not an error, just not enough real "
                      f"flip history accumulated so far.")
                continue
            capital_base = compute_capital_base(trades, capital_per_trade=DEFAULT_MARGIN_PER_LOT.get(index_name, 200000))
            metrics = compute_metrics(trades, capital_base)
            print_summary(metrics, excluded)
            if metrics:
                try:
                    path = write_pdf_report(trades, metrics)
                    print(f"\nFull PDF report written to: {path}")
                except ImportError:
                    print("\nPDF generation needs matplotlib and reportlab -- pip install matplotlib reportlab")
