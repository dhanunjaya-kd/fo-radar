"""
screener/check_oi_confirmation_score_bias.py

One-off diagnostic -- tests whether the OI-confirmation score bonus is
"rescuing" weaker technical setups into qualifying (Grade B/A), rather
than genuinely marking stronger setups. Built to answer a real question
raised by the Aug 2026 backtest reports: OI Confirmation "NEUTRAL"
trades showed ~8x the per-trade net P&L of "CONFIRMED" trades (Rs 3,988
vs Rs 481/trade), and Grade C (mostly NEUTRAL/NO_DATA) outright beat
Grade B (which needs a positive OI adjustment to reach) in the same
data. One explanation: CONFIRMED signals only needed a WEAKER base
setup to qualify at all, because the OI bonus (+20, or +25 with the PCR
kicker) pushed them over the score>=50 gate that a NEUTRAL signal had
to clear on RSI/Volume/ADX/alignment alone.

Reconstructs each signal's BASE score (before the OI adjustment) from
data ALREADY logged in signals_*.xlsx -- no new logging needed.
excel_logger.py's "Confidence" column IS the exact numeric total_score
(views.py's _build_all(): "confidence": f"{total_score}%"), and PCR is
logged separately. Exact inverse of that same function's OI/PCR scoring:

  oi_adjustment = 20 if OI Confirmation == "CONFIRMED" else 0
  oi_adjustment += 5 if (Action=="BUY" and PCR>1.0) or (Action=="SELL" and PCR<0.7) else 0
  base_score = Confidence(%) - oi_adjustment

This is an EXACT reconstruction, not an estimate -- every input it
needs (Action, PCR, OI Confirmation, Confidence) is already a real
logged column, so there's no approximation or missing data to guess at
for any row where Confidence is present.

Run as a standalone script against real signal logs:
    cd backend
    venv\\Scripts\\activate
    python -m screener.check_oi_confirmation_score_bias
"""
import glob
import os
import statistics

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

LOG_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "signal_logs")


def reconstruct_base_score(action, pcr, oi_confirmation, confidence_pct):
    """
    Exact inverse of views.py's _build_all() OI/PCR scoring -- see this
    module's own docstring for the full formula and why it's exact,
    not approximate. confidence_pct is the already-parsed numeric
    value (the "%" already stripped by the caller).
    """
    oi_adjustment = 20 if oi_confirmation == "CONFIRMED" else 0
    if pcr is not None:
        if action == "BUY" and pcr > 1.0:
            oi_adjustment += 5
        elif action == "SELL" and pcr < 0.7:
            oi_adjustment += 5
    return confidence_pct - oi_adjustment


def load_all_signal_rows():
    """
    Every logged signal row across every day found in LOG_DIR, pulling
    ONLY the columns this analysis needs. Deliberately a separate,
    narrower reader from backtest_signal_pnl.py's load_all_trades(),
    which doesn't currently capture Confidence/PCR at all -- adding
    them there for a one-off diagnostic isn't worth the risk to an
    already-tested pipeline. Returns a list of dicts; a row missing a
    usable Confidence value is skipped (nothing to reconstruct from),
    counted and reported, never silently dropped.
    """
    rows = []
    skipped_no_confidence = 0
    paths = sorted(glob.glob(os.path.join(LOG_DIR, "*", "signals_*.xlsx")))
    paths += sorted(glob.glob(os.path.join(LOG_DIR, "signals_*.xlsx")))  # old flat layout, pre-reorg
    for path in paths:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb["Signals"]
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            col = {name: i for i, name in enumerate(headers)}
            required = {"Action", "PCR", "OI Confirmation", "Confidence", "Grade", "Outcome"}
            if not required.issubset(col.keys()):
                print(f"  (skipping {os.path.basename(path)}: missing required column(s))")
                continue
            for raw in ws.iter_rows(min_row=2, values_only=True):
                action = raw[col["Action"]]
                oi_conf = raw[col["OI Confirmation"]]
                confidence_raw = raw[col["Confidence"]]
                pcr_raw = raw[col["PCR"]]
                grade = raw[col["Grade"]]
                outcome = raw[col["Outcome"]]
                if not action:
                    continue
                if not confidence_raw:
                    skipped_no_confidence += 1
                    continue
                try:
                    confidence_pct = float(str(confidence_raw).rstrip("%"))
                except (ValueError, TypeError):
                    skipped_no_confidence += 1
                    continue
                try:
                    pcr = float(pcr_raw) if pcr_raw not in (None, "") else None
                except (ValueError, TypeError):
                    pcr = None
                rows.append({
                    "action": action, "pcr": pcr, "oi_confirmation": oi_conf,
                    "confidence_pct": confidence_pct, "grade": grade, "outcome": outcome,
                    "source_file": os.path.basename(path),
                })
        except Exception as e:
            print(f"  (skipping {os.path.basename(path)}: {e})")
            continue
    if skipped_no_confidence:
        print(f"({skipped_no_confidence} row(s) skipped -- no usable Confidence value to reconstruct from)")
    return rows


def main():
    if load_workbook is None:
        print("openpyxl not installed -- pip install openpyxl --break-system-packages")
        return

    rows = load_all_signal_rows()
    if not rows:
        print("No signal rows found under", LOG_DIR)
        return

    for r in rows:
        r["base_score"] = reconstruct_base_score(r["action"], r["pcr"], r["oi_confirmation"], r["confidence_pct"])

    by_oi = {}
    for r in rows:
        by_oi.setdefault(r["oi_confirmation"], []).append(r)

    print(f"\nTotal signals loaded: {len(rows)}\n")
    header = f"{'OI Confirmation':<16} {'Count':>6} {'Avg Base Score':>16} {'Median Base':>13} {'Avg Confidence':>15}"
    print(header)
    print("-" * len(header))
    for state, group in sorted(by_oi.items(), key=lambda kv: -len(kv[1])):
        base_scores = [g["base_score"] for g in group]
        confidences = [g["confidence_pct"] for g in group]
        print(f"{str(state):<16} {len(group):>6} {statistics.mean(base_scores):>16.1f} "
              f"{statistics.median(base_scores):>13.1f} {statistics.mean(confidences):>15.1f}")

    print("\nHow to read this:")
    print("If CONFIRMED's Avg Base Score is meaningfully LOWER than NEUTRAL's,")
    print("that supports 'the OI bonus is rescuing weaker setups into qualifying'")
    print("-- the fix would be raising the base-score floor, not touching the OI")
    print("bonus itself. If the base scores are similar, that explanation doesn't")
    print("hold, and the real answer is more likely stale/mistimed OI data, or a")
    print("genuine but adverse crowding effect from one-sided OI writing.")


if __name__ == "__main__":
    main()
