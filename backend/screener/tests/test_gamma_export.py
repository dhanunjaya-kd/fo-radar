import io
from unittest.mock import patch

from django.test import TestCase
from openpyxl import load_workbook

import screener.views as views_module


class GammaExcelExportTests(TestCase):
    """
    Sep 30 2026: real tests for GammaStrategyExcelExportView. Each test
    patches the SAME module-level caches GammaStrategyView itself
    reads (confirmed by reading that view directly before writing
    this), then actually LOADS the returned bytes back with openpyxl
    to check real cell values -- not just that the endpoint returned
    200, which would miss a sheet silently containing the wrong data.
    """

    def _get_workbook(self):
        response = self.client.get("/api/gamma-strategy/export/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return load_workbook(io.BytesIO(response.content))

    def test_response_is_a_downloadable_attachment(self):
        with patch.object(views_module, "_gamma_watchlist_cache", {"resistance_watchlist": [], "support_watchlist": [], "updated_at": None, "symbols_with_zones_today": 0}), \
             patch.object(views_module, "_gamma_active_options_cache", {"items": [], "updated_at": None}), \
             patch.object(views_module, "_gamma_alerts_cache", {"items": []}):
            response = self.client.get("/api/gamma-strategy/export/")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])

    def test_all_five_expected_sheets_present(self):
        with patch.object(views_module, "_gamma_watchlist_cache", {"resistance_watchlist": [], "support_watchlist": [], "updated_at": None, "symbols_with_zones_today": 0}), \
             patch.object(views_module, "_gamma_active_options_cache", {"items": [], "updated_at": None}), \
             patch.object(views_module, "_gamma_alerts_cache", {"items": []}):
            wb = self._get_workbook()
        self.assertEqual(
            set(wb.sheetnames),
            {"Gamma Summary", "CE Candidates", "PE Candidates", "Options Resolver", "Microstructure Alerts"},
        )

    def test_ce_candidate_real_values_land_in_correct_sheet_and_columns(self):
        """The actual point of this whole feature: the exported file
        must contain the SAME data the page shows, not just a
        correctly-shaped empty template."""
        watchlist = {
            "resistance_watchlist": [
                {"symbol": "RELIANCE", "cmp": 2950.5, "zone_bottom": 2900, "zone_top": 2960,
                 "distance_pct": 0.32, "trend_aligned": True, "intraday_momentum": False, "status": "APPROACHING"},
            ],
            "support_watchlist": [],
            "updated_at": "2026-09-30T10:00:00",
            "symbols_with_zones_today": 5,
        }
        with patch.object(views_module, "_gamma_watchlist_cache", watchlist), \
             patch.object(views_module, "_gamma_active_options_cache", {"items": [], "updated_at": None}), \
             patch.object(views_module, "_gamma_alerts_cache", {"items": []}):
            wb = self._get_workbook()
        ws = wb["CE Candidates"]
        header = [c.value for c in ws[1]]
        row = [c.value for c in ws[2]]
        row_dict = dict(zip(header, row))
        self.assertEqual(row_dict["Symbol"], "RELIANCE")
        self.assertEqual(row_dict["CMP"], 2950.5)
        self.assertEqual(row_dict["50 EMA Aligned"], "Yes")
        self.assertEqual(row_dict["Intraday Momentum"], "No")
        # PE Candidates must stay empty -- this stock is a resistance
        # (CE) candidate only, confirming the two sheets aren't
        # accidentally merged or cross-contaminated.
        self.assertEqual(wb["PE Candidates"].max_row, 1)  # header row only

    def test_options_resolver_splits_ce_and_pe_correctly(self):
        options = {
            "items": [
                {"symbol": "TCS", "option_type": "CE", "strike": 4200, "tier": "GOLD", "ltp": 45.5,
                 "expiry": "30OCT", "dte": 12, "spread_pct": 1.2, "delta": 0.35, "convexity": 0.18,
                 "oi": 150000, "volume": 20000, "security_id": "ce1"},
                {"symbol": "INFY", "option_type": "PE", "strike": 1500, "tier": "SILVER", "ltp": 22.1,
                 "expiry": "30OCT", "dte": 12, "spread_pct": 0.9, "delta": -0.28, "convexity": 0.16,
                 "oi": 90000, "volume": 15000, "security_id": "pe1"},
            ],
            "updated_at": "2026-09-30T10:05:00",
        }
        with patch.object(views_module, "_gamma_watchlist_cache", {"resistance_watchlist": [], "support_watchlist": [], "updated_at": None, "symbols_with_zones_today": 0}), \
             patch.object(views_module, "_gamma_active_options_cache", options), \
             patch.object(views_module, "_gamma_alerts_cache", {"items": []}):
            wb = self._get_workbook()
        ws = wb["Options Resolver"]
        header = [c.value for c in ws[1]]
        rows = [dict(zip(header, [c.value for c in ws[r]])) for r in range(2, ws.max_row + 1)]
        symbols_and_types = {(r["Symbol"], r["Type"]) for r in rows}
        self.assertIn(("TCS", "CE"), symbols_and_types)
        self.assertIn(("INFY", "PE"), symbols_and_types)
        tcs_row = next(r for r in rows if r["Symbol"] == "TCS")
        self.assertEqual(tcs_row["Delta"], 0.35)
        self.assertEqual(tcs_row["OI"], 150000)

    def test_microstructure_alerts_real_values(self):
        alerts = [
            {"contract": "NIFTY24500CE", "status": "TARGET_1_HIT", "entry_price": 120.5,
             "stop_loss": 100.0, "target_1": 140.0, "target_2": 160.0, "timestamp_ist": "10:15:30", "alert_id": "a1"},
        ]
        with patch.object(views_module, "_gamma_watchlist_cache", {"resistance_watchlist": [], "support_watchlist": [], "updated_at": None, "symbols_with_zones_today": 0}), \
             patch.object(views_module, "_gamma_active_options_cache", {"items": [], "updated_at": None}), \
             patch.object(views_module, "_gamma_alerts_cache", {"items": alerts}):
            wb = self._get_workbook()
        ws = wb["Microstructure Alerts"]
        header = [c.value for c in ws[1]]
        row = dict(zip(header, [c.value for c in ws[2]]))
        self.assertEqual(row["Contract"], "NIFTY24500CE")
        self.assertEqual(row["Status"], "TARGET_1_HIT")
        self.assertEqual(row["Entry Price"], 120.5)

    def test_summary_sheet_reflects_real_counts_not_hardcoded(self):
        watchlist = {
            "resistance_watchlist": [{"symbol": "A"}, {"symbol": "B"}],
            "support_watchlist": [{"symbol": "C"}],
            "updated_at": "2026-09-30T10:00:00",
            "symbols_with_zones_today": 7,
        }
        with patch.object(views_module, "_gamma_watchlist_cache", watchlist), \
             patch.object(views_module, "_gamma_active_options_cache", {"items": [], "updated_at": None}), \
             patch.object(views_module, "_gamma_alerts_cache", {"items": []}):
            wb = self._get_workbook()
        ws = wb["Gamma Summary"]
        field_values = {row[0].value: row[1].value for row in ws.iter_rows(min_row=2)}
        self.assertEqual(field_values["Resistance Watchlist Count"], 2)
        self.assertEqual(field_values["Support Watchlist Count"], 1)
        self.assertEqual(field_values["Zone-Warmed Today"], 7)
        self.assertEqual(field_values["Options Resolver Status"], "WARMING_UP")  # no options items -- must reflect that honestly, not claim LIVE

    def test_empty_state_does_not_crash(self):
        """No fabricated rows when there's genuinely nothing yet --
        confirms the export degrades to empty sheets, not an error or
        invented sample data."""
        with patch.object(views_module, "_gamma_watchlist_cache", {"resistance_watchlist": [], "support_watchlist": [], "updated_at": None, "symbols_with_zones_today": 0}), \
             patch.object(views_module, "_gamma_active_options_cache", {"items": [], "updated_at": None}), \
             patch.object(views_module, "_gamma_alerts_cache", {"items": []}):
            wb = self._get_workbook()
        for name in ("CE Candidates", "PE Candidates", "Options Resolver", "Microstructure Alerts"):
            self.assertEqual(wb[name].max_row, 1)  # header only, no fabricated data rows
