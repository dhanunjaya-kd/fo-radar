"""Historical Next Day Watchlist research endpoints.

The live Tomorrow's Picks endpoint remains the source for the current
scan. These read-only endpoints expose the already-persisted Daily Picks
sheet from next_day_backtest.py so the UI can browse a specific scan date
without running another Fyers scan or changing ranking logic.
"""
from datetime import datetime
from io import BytesIO

from django.http import FileResponse, JsonResponse
from rest_framework.response import Response
from rest_framework.views import APIView

from .next_day_backtest import BACKTEST_FILE


def _parse_date(value):
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def _load_rows():
    if not BACKTEST_FILE:
        return []
    try:
        from openpyxl import load_workbook
        wb = load_workbook(BACKTEST_FILE, read_only=True, data_only=True)
        if "Daily Picks" not in wb.sheetnames:
            wb.close()
            return []
        ws = wb["Daily Picks"]
        rows = []
        headers = [cell.value for cell in ws[1]]
        positions = {str(v): i for i, v in enumerate(headers) if v is not None}
        for values in ws.iter_rows(min_row=2, values_only=True):
            scan_date = _parse_date(values[positions.get("Scan Date", 0)] if values else None)
            symbol = values[positions.get("Symbol", 2)] if len(values) > 2 else None
            if not scan_date or not symbol:
                continue
            def val(name, default=None):
                idx = positions.get(name)
                return values[idx] if idx is not None and idx < len(values) else default
            rows.append({
                "scan_date": scan_date.isoformat(),
                "rank": val("Rank"),
                "symbol": str(symbol),
                "sector": val("Sector"),
                "score": val("Score"),
                "eod_price": val("EOD Price"),
                "trend_status": val("Trend Status"),
                "volume_status": val("Volume Status"),
                "sector_strength": val("Sector Strength"),
            })
        wb.close()
        rows.sort(key=lambda r: (r["scan_date"], r.get("rank") or 9999), reverse=True)
        return rows
    except Exception as exc:
        print(f"[NextDayHistory] workbook read failed: {exc}")
        return []


class NextDayWatchlistDatesView(APIView):
    """GET /api/next-day-watchlist/dates/ -> available scan dates."""
    def get(self, request):
        rows = _load_rows()
        dates = sorted({row["scan_date"] for row in rows}, reverse=True)
        return Response({"dates": dates})


class NextDayWatchlistHistoryView(APIView):
    """GET /api/next-day-watchlist/history/?date=YYYY-MM-DD."""
    def get(self, request):
        date_str = request.GET.get("date")
        requested = _parse_date(date_str)
        if not requested:
            return Response({"error": "Valid date=YYYY-MM-DD is required."}, status=400)
        key = requested.isoformat()
        rows = [row for row in _load_rows() if row["scan_date"] == key]
        rows.sort(key=lambda r: r.get("rank") or 9999)
        return Response({
            "scan_date": key,
            "generated_at": f"{key}T15:40:00",
            "universe_scanned": None,
            "stocks_with_enough_data": None,
            "watchlist": rows,
            "historical": True,
        })


class NextDayWatchlistExportView(APIView):
    """Download one scan date as a small Excel research file."""
    def get(self, request):
        date_str = request.GET.get("date")
        requested = _parse_date(date_str)
        if not requested:
            return JsonResponse({"error": "Valid date=YYYY-MM-DD is required."}, status=400)
        key = requested.isoformat()
        rows = [row for row in _load_rows() if row["scan_date"] == key]
        if not rows:
            return JsonResponse({"error": f"No Tomorrow's Picks were recorded for {key}."}, status=404)

        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment

        wb = Workbook()
        ws = wb.active
        ws.title = "Tomorrow's Picks"
        headers = ["Scan Date", "Rank", "Symbol", "Sector", "Score", "EOD Price", "Trend Status", "Volume Status", "Sector Strength"]
        ws.append(headers)
        for row in rows:
            ws.append([
                row["scan_date"], row["rank"], row["symbol"], row["sector"], row["score"],
                row["eod_price"], row["trend_status"], row["volume_status"], row["sector_strength"],
            ])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        widths = [13, 8, 18, 18, 9, 14, 16, 22, 18]
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + idx)].width = width

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        filename = f"tomorrows_picks_{key}.xlsx"
        return FileResponse(
            buffer,
            as_attachment=True,
            filename=filename,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
