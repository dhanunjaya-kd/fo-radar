"""
Persistent lifecycle workbook for the Gamma Strategy block.

The exact option contract is the identity. An alert is written once; its
Entry/SL/T1/T2 remain fixed while status, hit times, current LTP, MFE/MAE,
realized R and carry-forward are updated in place across scan cycles and days.
"""
from __future__ import annotations
import glob, os, threading, time
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional

IST=timezone(timedelta(hours=5,30))
ROOT=Path(__file__).resolve().parents[1]
SIGNAL_LOGS=ROOT/"signal_logs"
TRACKER_PATH=SIGNAL_LOGS/"gamma_strategy_trade_tracking.xlsx"
COLUMNS=[
 "Contract","Status","Entry Price","Stop Loss","Target 1","Target 2","Timestamp (IST)",
 "T1 Hit At (IST)","T2 Hit At (IST)","SL Hit At (IST)","Trailing SL","Current LTP",
 "Last Checked At (IST)","MFE %","MAE %","Realized R","Carry Forward","Closed At (IST)",
 "Expiry","Security ID","Fyers Symbol","Symbol","Option Type","Strike","Alert ID","Trigger Candle","Trade Key"
]
ACTIVE={"ACTIVE","TARGET_1_HIT"}

try:
 from openpyxl import Workbook,load_workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 OPENPYXL_AVAILABLE=True
except ImportError:
 OPENPYXL_AVAILABLE=False

class _FileLock:
 def __init__(self,path,timeout=10.0): self.path=str(path)+".lock"; self.timeout=timeout; self.fd=None
 def __enter__(self):
  end=time.monotonic()+self.timeout
  while True:
   try: self.fd=os.open(self.path,os.O_CREAT|os.O_EXCL|os.O_RDWR); return self
   except FileExistsError:
    if time.monotonic()>=end:
     try: age=time.time()-os.path.getmtime(self.path)
     except OSError: age=None
     if age is not None and age>self.timeout:
      try: os.remove(self.path)
      except OSError: pass
      continue
     raise TimeoutError(f"Could not acquire lock on {self.path}")
    time.sleep(0.05)
 def __exit__(self,*args):
  if self.fd is not None: os.close(self.fd)
  try: os.remove(self.path)
  except OSError: pass
  return False

def _now(): return datetime.now(IST)
def _now_str(): return _now().strftime("%Y-%m-%d %H:%M:%S IST")
def _date(v):
 s=str(v or "").strip()
 if len(s)>=10:
  try: datetime.strptime(s[:10],"%Y-%m-%d"); return s[:10]
  except ValueError: pass
 return None
def _float(v,d=0.0):
 try: return float(v)
 except (TypeError,ValueError): return d
def contract_key(v):
 fs=str(v.get("fyers_symbol") or "").strip().upper()
 if fs: return fs
 c=str(v.get("contract") or "").strip().upper()
 if c: return c
 return "|".join([str(v.get("symbol") or "").strip().upper(),str(v.get("expiry") or "").strip().upper(),
                   str(v.get("strike") or "").strip().upper(),str(v.get("option_type") or "").strip().upper()])

def _read_daily(path):
 if not OPENPYXL_AVAILABLE or not path.exists(): return []
 try:
  wb=load_workbook(path,read_only=True,data_only=True)
  if "Gamma Signals" not in wb.sheetnames: wb.close(); return []
  ws=wb["Gamma Signals"]; h={str(c.value):i for i,c in enumerate(ws[1]) if c.value is not None}
  req={"Contract","Status","Entry Premium","Stop Loss","Target 1","Target 2"}
  if not req.issubset(h): wb.close(); return []
  out=[]
  for v in ws.iter_rows(min_row=2,values_only=True):
   if not v[h["Contract"]]: continue
   out.append({
    "contract":v[h["Contract"]],"status":v[h["Status"]] or "ACTIVE","entry_price":v[h["Entry Premium"]],
    "stop_loss":v[h["Stop Loss"]],"target_1":v[h["Target 1"]],"target_2":v[h["Target 2"]],
    "timestamp_ist":v[h["Timestamp"]] if "Timestamp" in h else None,
    "mfe_pct":v[h["MFE %"]] if "MFE %" in h else 0.0,"mae_pct":v[h["MAE %"]] if "MAE %" in h else 0.0,
    "realized_r":v[h["Realized R"]] if "Realized R" in h else 0.0,
    "closed_at_ist":v[h["Exit Time"]] if "Exit Time" in h else None,
    "expiry":v[h["Expiry"]] if "Expiry" in h else None,"symbol":v[h["Symbol"]] if "Symbol" in h else None,
    "option_type":v[h["Option Type"]] if "Option Type" in h else None,"strike":v[h["Strike"]] if "Strike" in h else None,
    "alert_id":v[h["Alert ID"]] if "Alert ID" in h else None,
    "trigger_candle":v[h["Trigger Candle"]] if "Trigger Candle" in h else None})
  wb.close(); return out
 except Exception as e:
  print(f"[GammaTradeTracker] daily bootstrap skipped: {e}"); return []

class GammaTradeTracker:
 def __init__(self,path:Optional[Path]=None):
  self.path=Path(os.environ.get("GAMMA_TRADE_TRACKER_XLSX",str(path or TRACKER_PATH)))
  SIGNAL_LOGS.mkdir(parents=True,exist_ok=True); self._lock=threading.RLock()
  self._ensure(); self._bootstrap()
 def _style(self,ws):
  for c in ws[1]:
   c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill(start_color="1F2937",end_color="1F2937",fill_type="solid"); c.alignment=Alignment(horizontal="center")
 def _ensure(self):
  if not OPENPYXL_AVAILABLE or self.path.exists(): return
  with self._lock,_FileLock(self.path):
   if self.path.exists(): return
   wb=Workbook(); ws=wb.active; ws.title="Gamma Trade Tracking"; ws.append(COLUMNS); self._style(ws); wb.save(self.path); wb.close()
 def _open(self):
  wb=load_workbook(self.path)
  if "Gamma Trade Tracking" in wb.sheetnames: return wb,wb["Gamma Trade Tracking"]
  ws=wb.create_sheet("Gamma Trade Tracking"); ws.append(COLUMNS); self._style(ws); return wb,ws
 @staticmethod
 def _headers(ws): return {str(c.value):i+1 for i,c in enumerate(ws[1]) if c.value is not None}
 def _read_all(self):
  wb,ws=self._open(); h=self._headers(ws); out={}
  for r in range(2,ws.max_row+1):
   key=ws.cell(r,h["Trade Key"]).value or ws.cell(r,h["Contract"]).value
   if not key: continue
   out[str(key).strip().upper()]={n:ws.cell(r,col).value for n,col in h.items()}|{"_row":r}
  wb.close(); return out
 def _write(self,ws,row,t):
  h=self._headers(ws); vals={
   "Contract":t.get("contract"),"Status":t.get("status","ACTIVE"),"Entry Price":t.get("entry_price"),
   "Stop Loss":t.get("stop_loss"),"Target 1":t.get("target_1"),"Target 2":t.get("target_2"),
   "Timestamp (IST)":t.get("timestamp_ist"),"T1 Hit At (IST)":t.get("t1_hit_at_ist"),"T2 Hit At (IST)":t.get("t2_hit_at_ist"),
   "SL Hit At (IST)":t.get("sl_hit_at_ist"),"Trailing SL":t.get("trailing_sl"),"Current LTP":t.get("current_ltp"),
   "Last Checked At (IST)":t.get("last_checked_at_ist"),"MFE %":t.get("mfe_pct"),"MAE %":t.get("mae_pct"),
   "Realized R":t.get("realized_r"),"Carry Forward":"YES" if str(t.get("carry_forward","")).upper() in ("YES","TRUE") or t.get("carry_forward") is True else "NO",
   "Closed At (IST)":t.get("closed_at_ist"),"Expiry":t.get("expiry"),"Security ID":t.get("security_id"),
   "Fyers Symbol":t.get("fyers_symbol"),"Symbol":t.get("symbol"),"Option Type":t.get("option_type"),
   "Strike":t.get("strike"),"Alert ID":t.get("alert_id"),"Trigger Candle":t.get("trigger_candle"),
   "Trade Key":t.get("_key") or contract_key(t)}
  for n,c in h.items():
   if n in vals: ws.cell(row=row,column=c,value=vals[n])
 def _save(self,wb,ws):
  self._style(ws); tmp=self.path.with_suffix(".tmp.xlsx"); wb.save(tmp); tmp.replace(self.path)
 def _bootstrap(self):
  files=sorted(glob.glob(str(SIGNAL_LOGS/"*"/"gamma_blast_*.xlsx")))
  if not files: return
  with self._lock,_FileLock(self.path):
   existing=self._read_all(); wb,ws=self._open(); changed=False
   for f in files:
    for raw in _read_daily(Path(f)):
     key=contract_key(raw)
     if not key or key in existing: continue
     raw["_key"]=key; raw["entry_price"]=_float(raw.get("entry_price")); raw["stop_loss"]=_float(raw.get("stop_loss"))
     raw["target_1"]=_float(raw.get("target_1")); raw["target_2"]=_float(raw.get("target_2"))
     raw["current_ltp"]=raw["entry_price"]; raw["trailing_sl"]=raw["stop_loss"]; raw["last_checked_at_ist"]=raw.get("timestamp_ist") or _now_str()
     raw["carry_forward"]=False; self._write(ws,ws.max_row+1,raw); existing[key]=raw; changed=True
   if changed: self._save(wb,ws)
   wb.close()
 def record_alert(self,alert,option_meta=None):
  merged=dict(option_meta or {}); merged.update(alert); key=contract_key(merged)
  if not key: return False
  with self._lock,_FileLock(self.path):
   if key in self._read_all(): return False
   wb,ws=self._open(); trade=dict(alert); trade["_key"]=key; trade["status"]="ACTIVE"; trade["current_ltp"]=alert.get("entry_price")
   trade["trailing_sl"]=alert.get("trailing_sl",alert.get("stop_loss")); trade["last_checked_at_ist"]=_now_str()
   trade["carry_forward"]=False; trade["closed_at_ist"]=None; trade["t1_hit_at_ist"]=None; trade["t2_hit_at_ist"]=None; trade["sl_hit_at_ist"]=None
   self._write(ws,ws.max_row+1,trade); self._save(wb,ws); wb.close(); return True
 def sync_alerts(self,alerts):
  if not OPENPYXL_AVAILABLE: return
  with self._lock,_FileLock(self.path):
   existing=self._read_all(); wb,ws=self._open(); today=_now().strftime("%Y-%m-%d"); changed=False
   for a in alerts or []:
    key=contract_key(a)
    if not key or key not in existing: continue
    row=existing[key]; status=str(a.get("status") or row.get("Status") or "ACTIVE").upper(); created=_date(row.get("Timestamp (IST)"))
    t={"contract":row.get("Contract"),"status":status,"entry_price":row.get("Entry Price"),"stop_loss":row.get("Stop Loss"),"target_1":row.get("Target 1"),"target_2":row.get("Target 2"),
       "timestamp_ist":row.get("Timestamp (IST)"),"t1_hit_at_ist":a.get("t1_hit_at_ist") or row.get("T1 Hit At (IST)"),"t2_hit_at_ist":a.get("t2_hit_at_ist") or row.get("T2 Hit At (IST)"),
       "sl_hit_at_ist":a.get("sl_hit_at_ist") or row.get("SL Hit At (IST)"),"trailing_sl":a.get("trailing_sl",row.get("Trailing SL")),
       "current_ltp":a.get("current_ltp") if a.get("current_ltp") is not None else (a.get("exit_price") if a.get("exit_price") is not None else row.get("Current LTP")),
       "last_checked_at_ist":_now_str(),"mfe_pct":a.get("mfe_pct",row.get("MFE %")),"mae_pct":a.get("mae_pct",row.get("MAE %")),
       "realized_r":a.get("realized_r",row.get("Realized R")),"carry_forward":bool(created and created<today and status in ACTIVE),
       "closed_at_ist":(a.get("exit_time_ist") if status in {"TARGET_2_HIT","STOPPED_OUT","TARGET_1_HIT_TRAILED","EXPIRED"} else row.get("Closed At (IST)")),"expiry":a.get("expiry") or row.get("Expiry"),
       "security_id":a.get("security_id") or row.get("Security ID"),"fyers_symbol":a.get("fyers_symbol") or row.get("Fyers Symbol"),
       "symbol":a.get("symbol") or row.get("Symbol"),"option_type":a.get("option_type") or row.get("Option Type"),"strike":a.get("strike") or row.get("Strike"),
       "alert_id":a.get("alert_id") or row.get("Alert ID"),"trigger_candle":a.get("trigger_candle") or row.get("Trigger Candle"),"_key":key}
    self._write(ws,row["_row"],t); changed=True
   if changed: self._save(wb,ws)
   wb.close()
 def get_active_contracts(self):
  with self._lock:
   rows=self._read_all()
   return [{"trade_key":k,"contract":r.get("Contract"),"fyers_symbol":r.get("Fyers Symbol"),"security_id":r.get("Security ID"),
            "symbol":r.get("Symbol"),"option_type":r.get("Option Type"),"strike":r.get("Strike"),"expiry":r.get("Expiry")}
           for k,r in rows.items() if str(r.get("Status") or "").upper() in ACTIVE]
 def is_known(self,option):
  k=contract_key(option)
  return bool(k and k in self._read_all())
 def is_active(self,option):
  k=contract_key(option); r=self._read_all().get(k) if k else None
  return bool(r and str(r.get("Status") or "").upper() in ACTIVE)
 def rebuild_open_alerts(self):
  with self._lock:
   rows=self._read_all(); out=[]
   for key,r in rows.items():
    if str(r.get("Status") or "").upper() not in ACTIVE: continue
    out.append({"alert_id":r.get("Alert ID") or f"PERSISTED_{key}","strategy":"Gamma_Blast_Options_strategy",
      "setup_type":"RESISTANCE_BREAKOUT_CE" if r.get("Option Type")=="CE" else "SUPPORT_BREAKDOWN_PE","horizon":"INTRADAY / SWING",
      "symbol":r.get("Symbol") or "","contract":r.get("Contract") or key,"security_id":r.get("Security ID") or key,
      "option_type":r.get("Option Type"),"strike":r.get("Strike"),"expiry":r.get("Expiry"),"lot_size":1,
      "entry_price":_float(r.get("Entry Price")),"stop_loss":_float(r.get("Stop Loss")),"trailing_sl":_float(r.get("Trailing SL"),_float(r.get("Stop Loss"))),
      "target_1":_float(r.get("Target 1")),"target_2":_float(r.get("Target 2")),"risk_reward":"1:1.6 to 1:2.8",
      "trigger_candle":r.get("Trigger Candle"),"status":r.get("Status"),"highest_ltp":_float(r.get("Current LTP"),_float(r.get("Entry Price"))),
      "lowest_ltp":_float(r.get("Current LTP"),_float(r.get("Entry Price"))),"mfe_pct":_float(r.get("MFE %")),"mae_pct":_float(r.get("MAE %")),
      "realized_r":_float(r.get("Realized R")),"exit_price":None,"exit_time_ist":None,"timestamp_ist":r.get("Timestamp (IST)"),
      "fyers_symbol":r.get("Fyers Symbol") or ""})
   return out
 def get_all_rows(self):
  with self._lock:
   rows=self._read_all(); out=[]
   for r in rows.values():
    out.append({"contract":r.get("Contract"),"status":r.get("Status"),"entry_price":r.get("Entry Price"),"stop_loss":r.get("Stop Loss"),
      "target_1":r.get("Target 1"),"target_2":r.get("Target 2"),"timestamp_ist":r.get("Timestamp (IST)"),
      "t1_hit_at_ist":r.get("T1 Hit At (IST)"),"t2_hit_at_ist":r.get("T2 Hit At (IST)"),"sl_hit_at_ist":r.get("SL Hit At (IST)"),
      "trailing_sl":r.get("Trailing SL"),"current_ltp":r.get("Current LTP"),"last_checked_at_ist":r.get("Last Checked At (IST)"),
      "mfe_pct":r.get("MFE %"),"mae_pct":r.get("MAE %"),"realized_r":r.get("Realized R"),"carry_forward":r.get("Carry Forward"),
      "closed_at_ist":r.get("Closed At (IST)"),"expiry":r.get("Expiry"),"security_id":r.get("Security ID"),"fyers_symbol":r.get("Fyers Symbol"),
      "symbol":r.get("Symbol"),"option_type":r.get("Option Type"),"strike":r.get("Strike"),"alert_id":r.get("Alert ID"),
      "trigger_candle":r.get("Trigger Candle"),"trade_key":r.get("Trade Key")})
   out.sort(key=lambda x:str(x.get("timestamp_ist") or "")); return out
